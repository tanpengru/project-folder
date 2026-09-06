# ============================================================
# PQAI.py
# NRF Parliamentary Questions AI Assistant
#
# Functions:
# 1. Parliamentary Question / Answer processing
# 2. NRF document processing
# 3. Text extraction
# 4. Text chunking
# 5. Semantic embeddings
# 6. Semantic search
# 7. Hybrid keyword + semantic search
# 8. Search index management
#
# IMPORTANT:
# This module is self-contained.
# It does NOT import main.py.
# ============================================================

import os
import re
import json
import zipfile
import logging
import tempfile
from functools import lru_cache

import numpy as np
import pandas as pd
import pdfplumber

from PyPDF2 import PdfReader
from docx import Document
from lxml import etree
from pdf2image import convert_from_path
import pytesseract

from sentence_transformers import SentenceTransformer


# ============================================================
# 📝 LOGGING
# ============================================================

logger = logging.getLogger(__name__)


# ============================================================
# 📁 GLOBAL CONSTANTS
# ============================================================

SEARCH_INDEX_FILE = "nrf_search_index.csv"
EMBEDDINGS_FILE = "nrf_embeddings.npy"

MODEL_NAME = "all-MiniLM-L6-v2"

DEFAULT_CHUNK_SIZE = 300
DEFAULT_CHUNK_OVERLAP = 50

EMBEDDING_DIMENSION = 384


SEARCH_COLUMNS = [
    "chunk_id",
    "source_type",
    "source_id",
    "title",
    "text",
    "embedding_text",
    "filepath",
    "metadata",
]


# ============================================================
# 🧠 EMBEDDING MODEL
# ============================================================

_embedding_model = None


def get_embedding_model():
    """
    Lazily load and cache the SentenceTransformer model.

    all-MiniLM-L6-v2 produces 384-dimensional embeddings and
    provides a good speed / quality trade-off for this internal
    semantic-search repository.
    """

    global _embedding_model

    if _embedding_model is None:

        logger.info(
            "Loading SentenceTransformer model: %s",
            MODEL_NAME
        )

        _embedding_model = SentenceTransformer(
            MODEL_NAME
        )

    return _embedding_model


# ============================================================
# 🔧 VALUE NORMALIZATION
# ============================================================

def safe_string(value):
    """
    Convert a value into a clean string.

    Prevents pandas NaN values from becoming the literal
    searchable text 'nan'.
    """

    if value is None:
        return ""

    try:

        if pd.isna(value):
            return ""

    except (TypeError, ValueError):
        pass

    return str(value).strip()


# ============================================================
# 🔧 TEXT CLEANING
# ============================================================

def clean_text(text: str) -> str:
    """
    Clean extracted document text while preserving paragraphs.
    """

    if text is None:
        return ""

    text = str(text)

    # Remove null characters.
    text = text.replace(
        "\x00",
        " "
    )

    # Normalize line endings.
    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    # Normalize spaces but preserve newlines.
    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    # Remove spaces around newlines.
    text = re.sub(
        r" *\n *",
        "\n",
        text
    )

    # Prevent huge blank sections.
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# ✂️ TEXT CHUNKING
# ============================================================

def chunk_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_CHUNK_OVERLAP
):
    """
    Split text into overlapping word-based chunks.

    chunk_size:
        Approximate number of words per chunk.

    overlap:
        Number of words shared between adjacent chunks.
    """

    text = clean_text(text)

    if not text:
        return []

    try:
        chunk_size = int(chunk_size)
    except (TypeError, ValueError):
        chunk_size = DEFAULT_CHUNK_SIZE

    try:
        overlap = int(overlap)
    except (TypeError, ValueError):
        overlap = DEFAULT_CHUNK_OVERLAP

    chunk_size = max(
        1,
        chunk_size
    )

    overlap = max(
        0,
        min(
            overlap,
            chunk_size - 1
        )
    )

    words = text.split()

    if not words:
        return []

    chunks = []

    start = 0

    while start < len(words):

        end = min(
            start + chunk_size,
            len(words)
        )

        chunk = " ".join(
            words[start:end]
        ).strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(words):
            break

        start = end - overlap

    return chunks


# ============================================================
# 📄 DOCX TEXTBOX EXTRACTION
# ============================================================

def extract_textboxes_from_docx(filepath):
    """
    Extract text contained inside Word textboxes and shapes.
    """

    text_chunks = []

    namespaces = {

        "w":
            "http://schemas.openxmlformats.org/"
            "wordprocessingml/2006/main",

        "wps":
            "http://schemas.microsoft.com/office/"
            "word/2010/wordprocessingShape",

        "v":
            "urn:schemas-microsoft-com:vml",
    }

    try:

        with zipfile.ZipFile(
            filepath,
            "r"
        ) as docx:

            for name in docx.namelist():

                if not name.startswith(
                    "word/"
                ):
                    continue

                if not name.endswith(
                    ".xml"
                ):
                    continue

                try:

                    xml_data = docx.read(
                        name
                    )

                    root = etree.fromstring(
                        xml_data
                    )

                except Exception:
                    continue

                # Standard Word textboxes
                for node in root.findall(
                    ".//w:txbxContent",
                    namespaces
                ):

                    text_chunks.extend(
                        node.itertext()
                    )

                # Word 2010 shapes
                for node in root.findall(
                    ".//wps:txbx",
                    namespaces
                ):

                    text_chunks.extend(
                        node.itertext()
                    )

                # Legacy VML textboxes
                for node in root.findall(
                    ".//v:textbox",
                    namespaces
                ):

                    text_chunks.extend(
                        node.itertext()
                    )

    except Exception as exc:

        logger.warning(
            "Could not extract DOCX textboxes from %s: %s",
            filepath,
            exc
        )

    return "\n".join(
        clean_text(value)
        for value in text_chunks
        if clean_text(value)
    )


# ============================================================
# 📕 PDF EXTRACTION
# ============================================================

def extract_text_from_pdf(filepath):
    """
    Extract PDF text using:

    1. pdfplumber
    2. PyPDF2 fallback
    3. OCR fallback
    """

    parts = []

    # --------------------------------------------------------
    # 1. pdfplumber
    # --------------------------------------------------------

    try:

        with pdfplumber.open(
            filepath
        ) as pdf:

            for page in pdf.pages:

                text = page.extract_text()

                if text and text.strip():

                    parts.append(
                        text
                    )

    except Exception as exc:

        logger.warning(
            "pdfplumber failed for %s: %s",
            filepath,
            exc
        )

    # --------------------------------------------------------
    # 2. PyPDF2
    # --------------------------------------------------------

    if not parts:

        try:

            reader = PdfReader(
                filepath
            )

            extracted = []

            for page in reader.pages:

                text = (
                    page.extract_text()
                    or ""
                )

                if text.strip():

                    extracted.append(
                        text
                    )

            if extracted:

                parts.extend(
                    extracted
                )

        except Exception as exc:

            logger.warning(
                "PyPDF2 failed for %s: %s",
                filepath,
                exc
            )

    # --------------------------------------------------------
    # 3. OCR
    # --------------------------------------------------------

    if not parts:

        try:

            images = convert_from_path(
                filepath,
                dpi=200
            )

            ocr_parts = []

            for image in images:

                text = (
                    pytesseract
                    .image_to_string(
                        image
                    )
                )

                if text.strip():

                    ocr_parts.append(
                        text
                    )

            if ocr_parts:

                parts.extend(
                    ocr_parts
                )

        except Exception as exc:

            logger.warning(
                "OCR failed for %s: %s",
                filepath,
                exc
            )

    return clean_text(
        "\n\n".join(parts)
    )


# ============================================================
# 📘 DOCX EXTRACTION
# ============================================================

def extract_text_from_docx(filepath):
    """
    Extract normal paragraphs, tables and textboxes from DOCX.
    """

    parts = []

    try:

        doc = Document(
            filepath
        )

        # ----------------------------------------------------
        # Paragraphs
        # ----------------------------------------------------

        for paragraph in doc.paragraphs:

            text = clean_text(
                paragraph.text
            )

            if text:

                parts.append(
                    text
                )

        # ----------------------------------------------------
        # Tables
        # ----------------------------------------------------

        for table in doc.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    value = clean_text(
                        cell.text
                    )

                    if value:

                        cells.append(
                            value
                        )

                if cells:

                    parts.append(
                        " | ".join(cells)
                    )

    except Exception as exc:

        logger.warning(
            "DOCX extraction failed for %s: %s",
            filepath,
            exc
        )

    # --------------------------------------------------------
    # Textboxes
    # --------------------------------------------------------

    textbox_text = (
        extract_textboxes_from_docx(
            filepath
        )
    )

    if textbox_text:

        parts.append(
            textbox_text
        )

    return clean_text(
        "\n".join(parts)
    )


# ============================================================
# 📊 EXCEL EXTRACTION
# ============================================================

def extract_text_from_excel(filepath):
    """
    Convert Excel sheets into searchable plain text.
    """

    parts = []

    try:

        excel_file = pd.ExcelFile(
            filepath
        )

    except Exception as exc:

        logger.warning(
            "Could not open Excel file %s: %s",
            filepath,
            exc
        )

        return ""

    for sheet_name in excel_file.sheet_names:

        try:

            df = pd.read_excel(
                filepath,
                sheet_name=sheet_name,
                dtype=str
            )

            df = df.fillna("")

            parts.append(
                f"=== SHEET: {sheet_name} ==="
            )

            if len(df.columns):

                parts.append(
                    " | ".join(
                        safe_string(column)
                        for column in df.columns
                    )
                )

            for _, row in df.iterrows():

                values = [
                    safe_string(value)
                    for value in row.tolist()
                    if safe_string(value)
                ]

                if values:

                    parts.append(
                        " | ".join(values)
                    )

        except Exception as exc:

            logger.warning(
                "Could not read Excel sheet %s "
                "from %s: %s",
                sheet_name,
                filepath,
                exc
            )

    return clean_text(
        "\n".join(parts)
    )


# ============================================================
# 📃 TXT EXTRACTION
# ============================================================

def extract_text_from_txt(filepath):

    try:

        with open(
            filepath,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            return clean_text(
                file.read()
            )

    except Exception as exc:

        logger.warning(
            "TXT extraction failed for %s: %s",
            filepath,
            exc
        )

        return ""


# ============================================================
# 📄 GENERAL DOCUMENT EXTRACTION
# ============================================================

@lru_cache(maxsize=128)
def _extract_raw_text_cached(
    filepath,
    modified_time
):
    """
    Internal cached document extraction.

    modified_time is included in the cache key so replacing
    a file invalidates its previous cached extraction.
    """

    del modified_time

    extension = os.path.splitext(
        filepath
    )[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(
            filepath
        )

    if extension == ".docx":

        return extract_text_from_docx(
            filepath
        )

    if extension in {
        ".xlsx",
        ".xls"
    }:

        return extract_text_from_excel(
            filepath
        )

    if extension == ".txt":

        return extract_text_from_txt(
            filepath
        )

    logger.warning(
        "Unsupported document type: %s",
        extension
    )

    return ""


def extract_raw_text(filepath):
    """
    Extract text from a repository document.

    This function deliberately lives inside PQAI.py rather
    than importing it from main.py. This avoids the previous
    main.py <-> PQAI.py circular import.
    """

    filepath = safe_string(
        filepath
    )

    if not filepath:

        return ""

    if not os.path.exists(
        filepath
    ):

        logger.warning(
            "Repository file does not exist: %s",
            filepath
        )

        return ""

    try:

        modified_time = os.path.getmtime(
            filepath
        )

    except OSError:

        modified_time = 0

    return _extract_raw_text_cached(
        filepath,
        modified_time
    )


# ============================================================
# 🏛️ PQ DISPLAY / EMBEDDING TEXT
# ============================================================

def build_pq_search_text(row):
    """
    Build semantic-search text for a Parliamentary Question.

    PQ question and answer remain together because a query may
    correspond to either side of the parliamentary exchange.
    """

    date = safe_string(
        row.get(
            "date",
            ""
        )
    )

    sitting = safe_string(
        row.get(
            "sitting",
            ""
        )
    )

    mp_name = safe_string(
        row.get(
            "mp_name",
            ""
        )
    )

    constituency = safe_string(
        row.get(
            "constituency",
            ""
        )
    )

    topic = safe_string(
        row.get(
            "topic",
            ""
        )
    )

    keywords = safe_string(
        row.get(
            "keywords",
            ""
        )
    )

    question = safe_string(
        row.get(
            "question",
            ""
        )
    )

    answer = safe_string(
        row.get(
            "answer",
            ""
        )
    )

    source = safe_string(
        row.get(
            "source",
            ""
        )
    )

    sections = []

    if date:
        sections.append(
            f"Date: {date}"
        )

    if sitting:
        sections.append(
            f"Parliamentary Sitting: {sitting}"
        )

    if mp_name:
        sections.append(
            f"Member of Parliament: {mp_name}"
        )

    if constituency:
        sections.append(
            f"Constituency: {constituency}"
        )

    if topic:
        sections.append(
            f"Topic: {topic}"
        )

    if keywords:
        sections.append(
            f"Keywords: {keywords}"
        )

    if question:
        sections.append(
            f"Question:\n{question}"
        )

    if answer:
        sections.append(
            f"NRF Answer:\n{answer}"
        )

    if source:
        sections.append(
            f"Source: {source}"
        )

    return clean_text(
        "\n\n".join(sections)
    )


# ============================================================
# 📚 DOCUMENT EMBEDDING CONTEXT
# ============================================================

def build_document_metadata_context(
    metadata
):
    """
    Build the metadata prefix supplied to the embedding model.

    This context is attached to EVERY document chunk rather
    than being inserted only once before chunking.
    """

    title = safe_string(
        metadata.get(
            "document_title",
            ""
        )
    )

    filename = safe_string(
        metadata.get(
            "filename",
            ""
        )
    )

    document_type = safe_string(
        metadata.get(
            "document_type",
            ""
        )
    )

    description = safe_string(
        metadata.get(
            "description",
            ""
        )
    )

    keywords = safe_string(
        metadata.get(
            "keywords",
            ""
        )
    )

    source = safe_string(
        metadata.get(
            "source",
            ""
        )
    )

    parts = []

    if title:

        parts.append(
            f"Document Title: {title}"
        )

    if filename:

        parts.append(
            f"Filename: {filename}"
        )

    if document_type:

        parts.append(
            f"Document Type: {document_type}"
        )

    if keywords:

        parts.append(
            f"Keywords: {keywords}"
        )

    if description:

        parts.append(
            f"Description: {description}"
        )

    if source:

        parts.append(
            f"Source: {source}"
        )

    return clean_text(
        "\n".join(parts)
    )


def build_document_search_text(
    text: str,
    metadata: dict
):
    """
    Backward-compatible helper.

    Returns document metadata plus content.
    """

    metadata_context = (
        build_document_metadata_context(
            metadata
        )
    )

    if metadata_context:

        return clean_text(
            metadata_context
            + "\n\n"
            + text
        )

    return clean_text(
        text
    )


# ============================================================
# 🏛️ CREATE PQ CHUNKS
# ============================================================

def create_pq_chunks(
    pq_df: pd.DataFrame
):
    """
    Convert Parliamentary Questions into searchable chunks.
    """

    chunks = []

    if (
        pq_df is None
        or pq_df.empty
    ):

        return chunks

    for _, row in pq_df.iterrows():

        pq_id = safe_string(
            row.get(
                "pq_id",
                ""
            )
        )

        if not pq_id:
            continue

        full_text = (
            build_pq_search_text(
                row
            )
        )

        if not full_text:
            continue

        pq_chunks = chunk_text(
            full_text,
            chunk_size=300,
            overlap=50
        )

        metadata = {

            "date":
                safe_string(
                    row.get(
                        "date",
                        ""
                    )
                ),

            "sitting":
                safe_string(
                    row.get(
                        "sitting",
                        ""
                    )
                ),

            "mp_name":
                safe_string(
                    row.get(
                        "mp_name",
                        ""
                    )
                ),

            "constituency":
                safe_string(
                    row.get(
                        "constituency",
                        ""
                    )
                ),

            "topic":
                safe_string(
                    row.get(
                        "topic",
                        ""
                    )
                ),

            "keywords":
                safe_string(
                    row.get(
                        "keywords",
                        ""
                    )
                ),

            "source":
                safe_string(
                    row.get(
                        "source",
                        ""
                    )
                )
        }

        date = metadata[
            "date"
        ]

        mp_name = metadata[
            "mp_name"
        ]

        title_parts = [
            "PQ"
        ]

        if date:
            title_parts.append(date)

        if mp_name:
            title_parts.append(mp_name)

        title = " — ".join(
            title_parts
        )

        for i, chunk in enumerate(
            pq_chunks
        ):

            # PQ display and embedding text are currently the
            # same because PQ metadata is useful in the result.
            embedding_text = chunk

            chunks.append(
                {

                    "chunk_id":
                        f"{pq_id}_{i}",

                    "source_type":
                        "Parliamentary Question",

                    "source_id":
                        pq_id,

                    "title":
                        title,

                    "text":
                        chunk,

                    "embedding_text":
                        embedding_text,

                    "filepath":
                        "",

                    "metadata":
                        json.dumps(
                            metadata,
                            ensure_ascii=False
                        )
                }
            )

    return chunks


# ============================================================
# 📚 CREATE DOCUMENT CHUNKS
# ============================================================

def create_document_chunks(
    document_df: pd.DataFrame
):
    """
    Convert NRF documents into searchable chunks.

    Display text:
        Only the relevant document passage.

    Embedding text:
        Document metadata + relevant passage.

    This provides semantic metadata context without repeatedly
    displaying metadata inside every result passage.
    """

    chunks = []

    if (
        document_df is None
        or document_df.empty
    ):

        return chunks

    for _, row in document_df.iterrows():

        document_id = safe_string(
            row.get(
                "document_id",
                ""
            )
        )

        filepath = safe_string(
            row.get(
                "filepath",
                ""
            )
        )

        if not document_id:

            logger.warning(
                "Skipping document with no document_id."
            )

            continue

        if (
            not filepath
            or not os.path.exists(filepath)
        ):

            logger.warning(
                "Skipping missing repository file: %s",
                filepath
            )

            continue

        # ----------------------------------------------------
        # Extract document text
        # ----------------------------------------------------

        text = extract_raw_text(
            filepath
        )

        if not text:

            logger.warning(
                "No searchable text extracted from: %s",
                filepath
            )

            continue

        # ----------------------------------------------------
        # Metadata
        # ----------------------------------------------------

        metadata = {

            "document_title":
                safe_string(
                    row.get(
                        "document_title",
                        ""
                    )
                ),

            "filename":
                safe_string(
                    row.get(
                        "filename",
                        ""
                    )
                ),

            "document_type":
                safe_string(
                    row.get(
                        "document_type",
                        ""
                    )
                ),

            "description":
                safe_string(
                    row.get(
                        "description",
                        ""
                    )
                ),

            "keywords":
                safe_string(
                    row.get(
                        "keywords",
                        ""
                    )
                ),

            "source":
                safe_string(
                    row.get(
                        "source",
                        ""
                    )
                )
        }

        metadata_context = (
            build_document_metadata_context(
                metadata
            )
        )

        # ----------------------------------------------------
        # Chunk ONLY the document body.
        # ----------------------------------------------------

        document_chunks = chunk_text(
            text,
            chunk_size=DEFAULT_CHUNK_SIZE,
            overlap=DEFAULT_CHUNK_OVERLAP
        )

        title = (
            metadata["document_title"]
            or metadata["filename"]
            or "Untitled Document"
        )

        for i, chunk in enumerate(
            document_chunks
        ):

            if metadata_context:

                embedding_text = clean_text(
                    metadata_context
                    + "\n\n"
                    + "Relevant Passage:\n"
                    + chunk
                )

            else:

                embedding_text = chunk

            chunks.append(
                {

                    "chunk_id":
                        f"{document_id}_{i}",

                    "source_type":
                        "NRF Document",

                    "source_id":
                        document_id,

                    "title":
                        title,

                    # Clean passage shown to user.
                    "text":
                        chunk,

                    # Metadata-enhanced text embedded by model.
                    "embedding_text":
                        embedding_text,

                    "filepath":
                        filepath,

                    "metadata":
                        json.dumps(
                            metadata,
                            ensure_ascii=False
                        )
                }
            )

    return chunks


# ============================================================
# 🔨 BUILD COMPLETE SEARCH INDEX
# ============================================================

def build_search_index(
    pq_df: pd.DataFrame,
    document_df: pd.DataFrame
):
    """
    Build the in-memory NRF semantic-search index.

    This function does not write files. Persistence happens
    only after embeddings have been generated successfully.
    """

    all_chunks = []

    # Parliamentary Questions
    all_chunks.extend(
        create_pq_chunks(
            pq_df
        )
    )

    # NRF documents
    all_chunks.extend(
        create_document_chunks(
            document_df
        )
    )

    index_df = pd.DataFrame(
        all_chunks,
        columns=SEARCH_COLUMNS
    )

    return index_df


# ============================================================
# 🧠 GENERATE EMBEDDINGS
# ============================================================

def generate_embeddings(
    index_df: pd.DataFrame
):
    """
    Generate normalized semantic embeddings.

    IMPORTANT:
    This function returns embeddings but does not save them.
    Persistence is handled by rebuild_search_index() so the
    CSV and embedding matrix stay synchronized.
    """

    if (
        index_df is None
        or index_df.empty
    ):

        return np.empty(
            (
                0,
                EMBEDDING_DIMENSION
            ),
            dtype=np.float32
        )

    model = get_embedding_model()

    # Prefer metadata-enhanced embedding text.
    if "embedding_text" in index_df.columns:

        texts = (
            index_df[
                "embedding_text"
            ]
            .fillna("")
            .astype(str)
            .tolist()
        )

    else:

        texts = (
            index_df[
                "text"
            ]
            .fillna("")
            .astype(str)
            .tolist()
        )

    embeddings = model.encode(
        texts,
        show_progress_bar=False,
        normalize_embeddings=True,
        convert_to_numpy=True
    )

    embeddings = np.asarray(
        embeddings,
        dtype=np.float32
    )

    if embeddings.ndim != 2:

        raise ValueError(
            "Embedding model returned an invalid "
            "embedding matrix."
        )

    if len(embeddings) != len(index_df):

        raise ValueError(
            "Number of embeddings does not match "
            "number of search-index rows."
        )

    return embeddings


# ============================================================
# 💾 SAFE INDEX PERSISTENCE
# ============================================================

def _atomic_save_csv(
    dataframe,
    destination
):
    """
    Write CSV through a temporary file before replacing the
    existing search index.
    """

    directory = (
        os.path.dirname(
            os.path.abspath(
                destination
            )
        )
        or "."
    )

    fd, temp_path = tempfile.mkstemp(
        prefix="nrf_index_",
        suffix=".csv",
        dir=directory
    )

    os.close(fd)

    try:

        dataframe.to_csv(
            temp_path,
            index=False,
            encoding="utf-8-sig"
        )

        os.replace(
            temp_path,
            destination
        )

    finally:

        if os.path.exists(
            temp_path
        ):

            try:
                os.remove(temp_path)
            except OSError:
                pass


def _atomic_save_embeddings(
    embeddings,
    destination
):
    """
    Write NPY embeddings through a temporary file.
    """

    directory = (
        os.path.dirname(
            os.path.abspath(
                destination
            )
        )
        or "."
    )

    fd, temp_path = tempfile.mkstemp(
        prefix="nrf_embeddings_",
        suffix=".npy",
        dir=directory
    )

    os.close(fd)

    try:

        np.save(
            temp_path,
            embeddings
        )

        os.replace(
            temp_path,
            destination
        )

    finally:

        if os.path.exists(
            temp_path
        ):

            try:
                os.remove(temp_path)
            except OSError:
                pass


# ============================================================
# 🔄 REBUILD COMPLETE INDEX
# ============================================================

def rebuild_search_index(
    pq_df: pd.DataFrame,
    document_df: pd.DataFrame
):
    """
    Rebuild the complete semantic search repository.

    Steps:
    1. Build index in memory.
    2. Generate all embeddings.
    3. Validate index / embedding alignment.
    4. Save both.

    Returns:
        The rebuilt search-index DataFrame.

    This return type intentionally matches main.py's need to
    report the number of searchable chunks using len(result).
    """

    logger.info(
        "Building NRF search index..."
    )

    index_df = build_search_index(
        pq_df,
        document_df
    )

    logger.info(
        "Generated %s searchable chunks.",
        len(index_df)
    )

    embeddings = generate_embeddings(
        index_df
    )

    if len(index_df) != len(embeddings):

        raise ValueError(
            "Search index and embedding matrix "
            "have different row counts."
        )

    # --------------------------------------------------------
    # Save only after everything succeeded.
    # --------------------------------------------------------

    _atomic_save_csv(
        index_df,
        SEARCH_INDEX_FILE
    )

    _atomic_save_embeddings(
        embeddings,
        EMBEDDINGS_FILE
    )

    logger.info(
        "NRF search index rebuilt successfully."
    )

    return index_df


# ============================================================
# 📂 LOAD SEARCH INDEX
# ============================================================

def load_search_index():
    """
    Load the persisted semantic-search index.
    """

    if not os.path.exists(
        SEARCH_INDEX_FILE
    ):

        return pd.DataFrame(
            columns=SEARCH_COLUMNS
        )

    try:

        df = pd.read_csv(
            SEARCH_INDEX_FILE,
            dtype=str,
            keep_default_na=False
        )

    except Exception as exc:

        logger.warning(
            "Could not load search index: %s",
            exc
        )

        return pd.DataFrame(
            columns=SEARCH_COLUMNS
        )

    # --------------------------------------------------------
    # Backward compatibility
    # --------------------------------------------------------

    for column in SEARCH_COLUMNS:

        if column not in df.columns:

            if (
                column == "embedding_text"
                and "text" in df.columns
            ):

                df[column] = df[
                    "text"
                ]

            else:

                df[column] = ""

    return df[
        SEARCH_COLUMNS
    ].reset_index(
        drop=True
    )


# ============================================================
# 🧠 LOAD EMBEDDINGS
# ============================================================

def load_embeddings():
    """
    Load persisted semantic embeddings.
    """

    if not os.path.exists(
        EMBEDDINGS_FILE
    ):

        return None

    try:

        embeddings = np.load(
            EMBEDDINGS_FILE,
            allow_pickle=False
        )

    except Exception as exc:

        logger.warning(
            "Could not load embeddings: %s",
            exc
        )

        return None

    if embeddings.ndim != 2:

        logger.warning(
            "Embedding file has invalid dimensions."
        )

        return None

    return embeddings


# ============================================================
# 🔍 VALIDATE INDEX
# ============================================================

def validate_search_index():
    """
    Check whether index CSV and embeddings are synchronized.

    Returns:
        (valid, message)
    """

    index_df = load_search_index()

    embeddings = load_embeddings()

    if embeddings is None:

        if index_df.empty:

            return (
                True,
                "Search index is empty."
            )

        return (
            False,
            "Embedding file is missing or invalid."
        )

    if len(index_df) != len(embeddings):

        return (
            False,
            (
                "Search index contains "
                f"{len(index_df)} rows but embedding "
                f"matrix contains {len(embeddings)} rows."
            )
        )

    return (
        True,
        (
            f"Search index is healthy with "
            f"{len(index_df)} searchable chunks."
        )
    )


# ============================================================
# 🔎 SEMANTIC SEARCH
# ============================================================

def semantic_search(
    query: str,
    top_k: int = 10,
    source_filter: str = "All Sources",
    min_similarity: float = 0.0
):
    """
    Perform semantic search across Parliamentary Questions
    and NRF documents.
    """

    query = clean_text(
        query
    )

    if not query:

        return pd.DataFrame()

    try:
        top_k = max(
            1,
            int(top_k)
        )
    except (TypeError, ValueError):
        top_k = 10

    try:
        min_similarity = float(
            min_similarity
        )
    except (TypeError, ValueError):
        min_similarity = 0.0

    # --------------------------------------------------------
    # Load index
    # --------------------------------------------------------

    index_df = load_search_index()

    if index_df.empty:

        return pd.DataFrame()

    embeddings = load_embeddings()

    if embeddings is None:

        logger.warning(
            "Semantic search unavailable: "
            "embedding file missing."
        )

        return pd.DataFrame()

    # --------------------------------------------------------
    # Critical consistency check
    # --------------------------------------------------------

    if len(index_df) != len(embeddings):

        logger.error(
            "Search index / embeddings mismatch: "
            "%s index rows vs %s embeddings.",
            len(index_df),
            len(embeddings)
        )

        return pd.DataFrame()

    # --------------------------------------------------------
    # Source filtering
    # --------------------------------------------------------

    if source_filter == "All Sources":

        selected_positions = np.arange(
            len(index_df)
        )

    else:

        selected_positions = np.flatnonzero(
            (
                index_df[
                    "source_type"
                ]
                .astype(str)
                .to_numpy()
                == source_filter
            )
        )

    if len(selected_positions) == 0:

        return pd.DataFrame()

    working_df = (
        index_df
        .iloc[
            selected_positions
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    selected_embeddings = embeddings[
        selected_positions
    ]

    # --------------------------------------------------------
    # Query embedding
    # --------------------------------------------------------

    model = get_embedding_model()

    query_embedding = model.encode(
        query,
        normalize_embeddings=True,
        convert_to_numpy=True
    )

    query_embedding = np.asarray(
        query_embedding,
        dtype=np.float32
    ).reshape(-1)

    # --------------------------------------------------------
    # Similarity
    #
    # Stored vectors and query vector are normalized.
    # Dot product therefore equals cosine similarity.
    # --------------------------------------------------------

    scores = (
        selected_embeddings
        @ query_embedding
    )

    working_df[
        "similarity"
    ] = scores.astype(float)

    # --------------------------------------------------------
    # Minimum score
    # --------------------------------------------------------

    working_df = working_df[
        working_df[
            "similarity"
        ]
        >= min_similarity
    ]

    if working_df.empty:

        return working_df.reset_index(
            drop=True
        )

    # --------------------------------------------------------
    # Sort
    # --------------------------------------------------------

    working_df = working_df.sort_values(
        "similarity",
        ascending=False
    )

    return (
        working_df
        .head(top_k)
        .reset_index(
            drop=True
        )
    )


# ============================================================
# 🏛️ SEARCH ONLY PARLIAMENTARY QUESTIONS
# ============================================================

def search_parliamentary_questions(
    query: str,
    top_k: int = 10,
    min_similarity: float = 0.0
):

    return semantic_search(
        query=query,
        top_k=top_k,
        source_filter="Parliamentary Question",
        min_similarity=min_similarity
    )


# ============================================================
# 📚 SEARCH ONLY NRF DOCUMENTS
# ============================================================

def search_nrf_documents(
    query: str,
    top_k: int = 10,
    min_similarity: float = 0.0
):

    return semantic_search(
        query=query,
        top_k=top_k,
        source_filter="NRF Document",
        min_similarity=min_similarity
    )


# ============================================================
# 🏛️ FIND SIMILAR PARLIAMENTARY QUESTIONS
# ============================================================

def find_similar_pqs(
    question: str,
    top_k: int = 10
):
    """
    Find historical PQs similar to a new Parliamentary
    Question.
    """

    return search_parliamentary_questions(
        query=question,
        top_k=top_k,
        min_similarity=0.0
    )


# ============================================================
# 📊 SEARCH STATISTICS
# ============================================================

def get_search_statistics():

    index_df = load_search_index()

    if index_df.empty:

        return {

            "total_chunks":
                0,

            "pq_chunks":
                0,

            "document_chunks":
                0
        }

    source_types = (
        index_df[
            "source_type"
        ]
        .fillna("")
        .astype(str)
    )

    pq_count = int(
        (
            source_types
            == "Parliamentary Question"
        ).sum()
    )

    document_count = int(
        (
            source_types
            == "NRF Document"
        ).sum()
    )

    return {

        "total_chunks":
            len(index_df),

        "pq_chunks":
            pq_count,

        "document_chunks":
            document_count
    }


# ============================================================
# 🧹 DELETE SOURCE FROM SEARCH INDEX
# ============================================================

def remove_source_from_index(
    source_type: str,
    source_id: str
):
    """
    Remove all chunks for one source.

    NOTE:
    main.py currently rebuilds the complete repository after
    deletion, which is safer. This function remains available
    for other callers.
    """

    index_df = load_search_index()

    if index_df.empty:

        return index_df

    source_type = safe_string(
        source_type
    )

    source_id = safe_string(
        source_id
    )

    mask = ~(
        (
            index_df[
                "source_type"
            ]
            == source_type
        )
        &
        (
            index_df[
                "source_id"
            ]
            == source_id
        )
    )

    remaining_df = (
        index_df[
            mask
        ]
        .reset_index(
            drop=True
        )
    )

    embeddings = generate_embeddings(
        remaining_df
    )

    _atomic_save_csv(
        remaining_df,
        SEARCH_INDEX_FILE
    )

    _atomic_save_embeddings(
        embeddings,
        EMBEDDINGS_FILE
    )

    return remaining_df


# ============================================================
# 🔍 EXACT TEXT SEARCH
# ============================================================

def exact_text_search(
    query: str,
    source_filter: str = "All Sources"
):
    """
    Traditional case-insensitive exact phrase search.

    Searches:
    - visible passage
    - embedding metadata context
    - title

    This is useful for acronyms, programme names, people,
    years and exact policy terminology.
    """

    query = clean_text(
        query
    )

    if not query:

        return pd.DataFrame()

    index_df = load_search_index()

    if index_df.empty:

        return pd.DataFrame()

    if source_filter != "All Sources":

        index_df = index_df[
            index_df[
                "source_type"
            ]
            == source_filter
        ].copy()

    if index_df.empty:

        return pd.DataFrame()

    q = query.lower()

    searchable = (
        index_df[
            "title"
        ]
        .fillna("")
        .astype(str)
        + "\n"
        +
        index_df[
            "text"
        ]
        .fillna("")
        .astype(str)
        + "\n"
        +
        index_df[
            "embedding_text"
        ]
        .fillna("")
        .astype(str)
    ).str.lower()

    mask = searchable.str.contains(
        re.escape(q),
        regex=True,
        na=False
    )

    results = index_df[
        mask
    ].copy()

    if results.empty:

        return results.reset_index(
            drop=True
        )

    results[
        "exact_match"
    ] = True

    return results.reset_index(
        drop=True
    )


# ============================================================
# 🔀 HYBRID SEARCH
# ============================================================

def hybrid_search(
    query: str,
    top_k: int = 10,
    source_filter: str = "All Sources",
    min_similarity: float = 0.0,
    exact_match_boost: float = 0.10
):
    """
    Combine semantic similarity with exact phrase matching.

    final_score =
        semantic similarity
        + exact phrase boost

    Exact matching is especially useful for:
    - programme names
    - acronyms
    - MP names
    - policy terminology
    - years
    """

    query = clean_text(
        query
    )

    if not query:

        return pd.DataFrame()

    try:

        top_k = max(
            1,
            int(top_k)
        )

    except (TypeError, ValueError):

        top_k = 10

    try:

        exact_match_boost = float(
            exact_match_boost
        )

    except (TypeError, ValueError):

        exact_match_boost = 0.10

    # --------------------------------------------------------
    # Retrieve broader semantic candidate pool
    # --------------------------------------------------------

    semantic_results = semantic_search(
        query=query,
        top_k=max(
            top_k * 3,
            30
        ),
        source_filter=source_filter,
        min_similarity=min_similarity
    )

    exact_results = exact_text_search(
        query=query,
        source_filter=source_filter
    )

    if (
        semantic_results.empty
        and exact_results.empty
    ):

        return pd.DataFrame()

    # --------------------------------------------------------
    # Prepare exact results
    # --------------------------------------------------------

    if not exact_results.empty:

        exact_results = (
            exact_results.copy()
        )

        if "similarity" not in exact_results.columns:

            exact_results[
                "similarity"
            ] = 0.0

    # --------------------------------------------------------
    # Combine
    # --------------------------------------------------------

    combined = pd.concat(
        [
            semantic_results,
            exact_results
        ],
        ignore_index=True,
        sort=False
    )

    combined = combined.drop_duplicates(
        subset=[
            "chunk_id"
        ],
        keep="first"
    )

    # --------------------------------------------------------
    # Identify exact matches
    # --------------------------------------------------------

    if exact_results.empty:

        exact_ids = set()

    else:

        exact_ids = set(
            exact_results[
                "chunk_id"
            ]
            .astype(str)
            .tolist()
        )

    combined[
        "exact_match"
    ] = (
        combined[
            "chunk_id"
        ]
        .astype(str)
        .isin(
            exact_ids
        )
    )

    combined[
        "similarity"
    ] = pd.to_numeric(
        combined[
            "similarity"
        ],
        errors="coerce"
    ).fillna(0.0)

    combined[
        "final_score"
    ] = (
        combined[
            "similarity"
        ]
        +
        (
            combined[
                "exact_match"
            ]
            .astype(float)
            * exact_match_boost
        )
    )

    combined = combined.sort_values(
        [
            "final_score",
            "similarity"
        ],
        ascending=[
            False,
            False
        ]
    )

    return (
        combined
        .head(top_k)
        .reset_index(
            drop=True
        )
    )


# ============================================================
# 🧾 FORMAT SEARCH RESULT
# ============================================================

def format_search_result(
    row
):
    """
    Convert a search-result row into a structured dictionary.
    """

    metadata = {}

    raw_metadata = row.get(
        "metadata",
        "{}"
    )

    if isinstance(
        raw_metadata,
        dict
    ):

        metadata = raw_metadata

    else:

        try:

            metadata = json.loads(
                safe_string(
                    raw_metadata
                )
                or "{}"
            )

        except Exception:

            metadata = {}

    try:

        relevance = float(
            row.get(
                "similarity",
                0
            )
        )

    except (TypeError, ValueError):

        relevance = 0.0

    result = {

        "source_type":
            safe_string(
                row.get(
                    "source_type",
                    ""
                )
            ),

        "source_id":
            safe_string(
                row.get(
                    "source_id",
                    ""
                )
            ),

        "title":
            safe_string(
                row.get(
                    "title",
                    ""
                )
            ),

        "relevance":
            relevance,

        "passage":
            safe_string(
                row.get(
                    "text",
                    ""
                )
            ),

        "filepath":
            safe_string(
                row.get(
                    "filepath",
                    ""
                )
            ),

        "metadata":
            metadata
    }

    if "final_score" in row:

        try:

            result[
                "final_score"
            ] = float(
                row.get(
                    "final_score",
                    relevance
                )
            )

        except (TypeError, ValueError):

            result[
                "final_score"
            ] = relevance

    return result


# ============================================================
# 📋 BUILD CONTEXT FOR AI
# ============================================================

def build_llm_context(
    search_results: pd.DataFrame,
    max_results: int = 10
):
    """
    Convert retrieved evidence into compact source-labelled
    context suitable for a later LLM prompt.
    """

    if (
        search_results is None
        or search_results.empty
    ):

        return ""

    try:

        max_results = max(
            1,
            int(max_results)
        )

    except (TypeError, ValueError):

        max_results = 10

    context_parts = []

    for result_number, (_, row) in enumerate(
        search_results
        .head(max_results)
        .iterrows(),
        start=1
    ):

        source_type = safe_string(
            row.get(
                "source_type",
                ""
            )
        )

        source_id = safe_string(
            row.get(
                "source_id",
                ""
            )
        )

        title = safe_string(
            row.get(
                "title",
                ""
            )
        )

        text = safe_string(
            row.get(
                "text",
                ""
            )
        )

        try:

            similarity = float(
                row.get(
                    "similarity",
                    0
                )
            )

        except (TypeError, ValueError):

            similarity = 0.0

        raw_metadata = row.get(
            "metadata",
            "{}"
        )

        if isinstance(
            raw_metadata,
            dict
        ):

            metadata = raw_metadata

        else:

            try:

                metadata = json.loads(
                    safe_string(
                        raw_metadata
                    )
                    or "{}"
                )

            except Exception:

                metadata = {}

        context_parts.append(
            f"""
SOURCE {result_number}

SOURCE TYPE:
{source_type}

SOURCE ID:
{source_id}

TITLE:
{title}

RELEVANCE:
{similarity:.1%}

SOURCE METADATA:
{json.dumps(
    metadata,
    ensure_ascii=False
)}

RELEVANT PASSAGE:
{text}

----------------------------------------
""".strip()
        )

    return "\n\n".join(
        context_parts
    )