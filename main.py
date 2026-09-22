# ============================================================
# main.py — NRF Parliamentary Questions AI Assistant
#
# Functions:
# 1. Store and manage past NRF Parliamentary Questions & Answers
# 2. Upload and manage NRF-related PDF, DOCX, Excel and TXT documents
# 3. AI chatbot grounded in:
#    - Parliamentary Questions / NRF answers
#    - Uploaded NRF documents
#
# Citation rules:
# - One unique Parliamentary Question = one [PQ#]
# - One unique uploaded NRF document = one [DOC#]
# - Multiple chunks from the same source use the SAME citation
#
# Search engine:
# helper_functions/PQAI.py
#
# AI:
# OpenAI Responses API
# ============================================================

import os
import re
import zipfile
import hashlib
import json
import importlib
import logging
import io
from collections import Counter
from datetime import datetime

import numpy as np
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin

try:
    from sklearn.linear_model import LinearRegression
except ImportError:
    LinearRegression = None

import streamlit as st
import pandas as pd

import pdfplumber
from PyPDF2 import PdfReader
from docx import Document
from lxml import etree
from pdf2image import convert_from_path
import pytesseract

from openai import OpenAI


# ============================================================
# 🔐 AUTH
# ============================================================

from helper_functions.utility import check_password


# ============================================================
# 🧠 NRF PQ AI SEARCH ENGINE
# ============================================================

try:
    from helper_functions import PQAI
except ImportError:
    PQAI = importlib.import_module("helper_functions.PQAI")


logger = logging.getLogger(__name__)


# ============================================================
# ⚙️ STREAMLIT CONFIG
# ============================================================

st.set_page_config(
    layout="wide",
    page_title="NRF Parliamentary Questions AI Assistant",
    page_icon="🏛️"
)

st.title("🏛️ NRF Parliamentary Questions AI Assistant")

if not check_password():
    st.stop()


# ============================================================
# 📁 GLOBAL CONSTANTS
# ============================================================

PQ_FOLDER = "nrf_pq_repo"
DOCUMENT_FOLDER = "nrf_documents"

PQ_METADATA_FILE = "nrf_pq_metadata.csv"
DOCUMENT_METADATA_FILE = "nrf_documents_metadata.csv"

os.makedirs(PQ_FOLDER, exist_ok=True)
os.makedirs(DOCUMENT_FOLDER, exist_ok=True)


# ============================================================
# 📊 METADATA SCHEMAS
# ============================================================

PQ_COLUMNS = [
    "pq_id",
    "date",
    "sitting",
    "mp_name",
    "constituency",
    "question",
    "answer",
    "topic",
    "keywords",
    "source",
    "uploaded_at",
]

DOCUMENT_COLUMNS = [
    "document_id",
    "filename",
    "filepath",
    "document_title",
    "document_type",
    "description",
    "keywords",
    "source",
    "uploaded_at",
]


# ============================================================
# 📚 GENERIC DATA HELPERS
# ============================================================

def load_csv(path, columns):
    """Load CSV while ensuring all expected columns exist."""

    if os.path.exists(path):
        try:
            df = pd.read_csv(
                path,
                dtype=str,
                keep_default_na=False
            )
        except Exception as exc:
            logger.exception(
                "Failed to load CSV %s: %s",
                path,
                exc
            )
            df = pd.DataFrame(columns=columns)
    else:
        df = pd.DataFrame(columns=columns)

    for col in columns:
        if col not in df.columns:
            df[col] = ""

    return df[columns]


def save_csv(df, path):
    """Save dataframe to CSV."""

    df.to_csv(
        path,
        index=False,
        encoding="utf-8-sig"
    )


def generate_id(prefix):
    """Generate short unique ID."""

    timestamp = datetime.now().isoformat()
    raw = f"{prefix}_{timestamp}"

    return hashlib.md5(
        raw.encode("utf-8")
    ).hexdigest()[:12]


def ensure_unique_path(folder, filename):
    """Prevent uploaded files from overwriting existing files."""

    base, ext = os.path.splitext(filename)

    path = os.path.join(
        folder,
        filename
    )

    counter = 1

    while os.path.exists(path):
        path = os.path.join(
            folder,
            f"{base} ({counter}){ext}"
        )
        counter += 1

    return path


# ============================================================
# 📄 DOCX TEXTBOX EXTRACTION
# ============================================================

def extract_textboxes_from_docx(filepath):

    text_chunks = []

    ns = {
        "w":
            "http://schemas.openxmlformats.org/"
            "wordprocessingml/2006/main",

        "wps":
            "http://schemas.microsoft.com/office/word/2010/"
            "wordprocessingShape",

        "v":
            "urn:schemas-microsoft-com:vml",
    }

    try:
        with zipfile.ZipFile(
            filepath,
            "r"
        ) as docx:

            for name in docx.namelist():

                if not name.startswith("word/"):
                    continue

                if not name.endswith(".xml"):
                    continue

                try:
                    xml_data = docx.read(name)
                    root = etree.fromstring(xml_data)
                except Exception:
                    continue

                for node in root.findall(
                    ".//w:txbxContent",
                    ns
                ):
                    text_chunks.extend(
                        node.itertext()
                    )

                for node in root.findall(
                    ".//wps:txbx",
                    ns
                ):
                    text_chunks.extend(
                        node.itertext()
                    )

                for node in root.findall(
                    ".//v:textbox",
                    ns
                ):
                    text_chunks.extend(
                        node.itertext()
                    )

    except Exception as exc:
        logger.exception(
            "Failed to extract DOCX textboxes from %s: %s",
            filepath,
            exc
        )

    return "\n".join(
        t.strip()
        for t in text_chunks
        if t.strip()
    )


# ============================================================
# 📕 PDF EXTRACTION
# ============================================================

def extract_text_from_pdf(filepath):

    parts = []

    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                text = page.extract_text()

                if text and text.strip():
                    parts.append(text)

    except Exception as exc:
        logger.warning(
            "pdfplumber failed for %s: %s",
            filepath,
            exc
        )

    try:
        reader = PdfReader(filepath)

        text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

        if text.strip() and not parts:
            parts.append(text)

    except Exception as exc:
        logger.warning(
            "PyPDF2 failed for %s: %s",
            filepath,
            exc
        )

    try:
        if not parts:

            images = convert_from_path(
                filepath,
                dpi=200
            )

            ocr_parts = []

            for image in images:
                text = pytesseract.image_to_string(
                    image
                )

                if text.strip():
                    ocr_parts.append(text)

            if ocr_parts:
                parts.append(
                    "\n".join(ocr_parts)
                )

    except Exception as exc:
        logger.warning(
            "OCR fallback failed for %s: %s",
            filepath,
            exc
        )

    return "\n\n".join(parts)


# ============================================================
# 📘 DOCX EXTRACTION
# ============================================================

def extract_text_from_docx(filepath):

    parts = []

    try:
        doc = Document(filepath)

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in doc.tables:
            for row in table.rows:

                cells = []

                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    parts.append(
                        " | ".join(cells)
                    )

    except Exception as exc:
        logger.warning(
            "DOCX extraction failed for %s: %s",
            filepath,
            exc
        )

    try:
        textbox_text = extract_textboxes_from_docx(
            filepath
        )

        if textbox_text.strip():
            parts.append(textbox_text)

    except Exception as exc:
        logger.warning(
            "DOCX textbox extraction failed for %s: %s",
            filepath,
            exc
        )

    return "\n".join(parts)


# ============================================================
# 📊 EXCEL EXTRACTION
# ============================================================

def extract_text_from_excel(filepath):

    parts = []

    try:
        excel_file = pd.ExcelFile(filepath)

        for sheet_name in excel_file.sheet_names:

            try:
                df = pd.read_excel(
                    filepath,
                    sheet_name=sheet_name,
                    dtype=str
                )

                df = df.fillna("")

                parts.append(
                    f"\n=== SHEET: {sheet_name} ==="
                )

                parts.append(
                    " | ".join(
                        str(col)
                        for col in df.columns
                    )
                )

                for _, row in df.iterrows():

                    values = [
                        str(value).strip()
                        for value in row.tolist()
                        if str(value).strip()
                    ]

                    if values:
                        parts.append(
                            " | ".join(values)
                        )

            except Exception as exc:
                logger.warning(
                    "Failed to read Excel sheet %s in %s: %s",
                    sheet_name,
                    filepath,
                    exc
                )

    except Exception as exc:
        logger.warning(
            "Excel extraction failed for %s: %s",
            filepath,
            exc
        )

    return "\n".join(parts)


# ============================================================
# 📄 GENERAL FILE EXTRACTION
# ============================================================

def extract_raw_text(filepath):

    if not filepath:
        return ""

    if not os.path.exists(filepath):
        return ""

    extension = os.path.splitext(
        filepath
    )[1].lower()

    if extension == ".pdf":
        return extract_text_from_pdf(filepath)

    elif extension == ".docx":
        return extract_text_from_docx(filepath)

    elif extension in [
        ".xlsx",
        ".xls"
    ]:
        return extract_text_from_excel(filepath)

    elif extension == ".txt":

        try:
            with open(
                filepath,
                "r",
                encoding="utf-8",
                errors="ignore"
            ) as f:
                return f.read()

        except Exception as exc:
            logger.warning(
                "TXT extraction failed for %s: %s",
                filepath,
                exc
            )

            return ""

    return ""


# ============================================================
# 🔄 SAFE SEARCH INDEX REBUILD
# ============================================================

def rebuild_search_index_safe(
    pq_dataframe,
    document_dataframe
):
    """
    Rebuild semantic search index safely.
    """

    try:

        rebuilt = PQAI.rebuild_search_index(
            pq_dataframe,
            document_dataframe
        )

        if isinstance(rebuilt, pd.DataFrame):
            return True, rebuilt

        if (
            isinstance(rebuilt, tuple)
            and len(rebuilt) >= 1
            and isinstance(rebuilt[0], pd.DataFrame)
        ):
            return True, rebuilt[0]

        logger.warning(
            "PQAI.rebuild_search_index returned "
            "unexpected type: %s",
            type(rebuilt).__name__
        )

        return False, pd.DataFrame()

    except Exception as exc:

        logger.exception(
            "Search index rebuild failed: %s",
            exc
        )

        return False, pd.DataFrame()


# ============================================================
# 🤖 OPENAI CLIENT
# ============================================================

def get_openai_api_key():

    try:
        if "OPENAI_API_KEY" in st.secrets:
            return st.secrets["OPENAI_API_KEY"]
    except Exception:
        pass

    return os.getenv(
        "OPENAI_API_KEY",
        ""
    )


def get_openai_model():

    try:
        if "OPENAI_MODEL" in st.secrets:
            return st.secrets["OPENAI_MODEL"]
    except Exception:
        pass

    return os.getenv(
        "OPENAI_MODEL",
        "gpt-4o-mini"
    )


def get_openai_client():

    api_key = get_openai_api_key()

    if not api_key:
        return None

    return OpenAI(
        api_key=api_key
    )


# ============================================================
# 🤖 UNIQUE SOURCE / CITATION HELPERS
# ============================================================

def _parse_source_metadata(metadata):
    """
    Convert metadata into a dictionary.
    """

    if isinstance(metadata, dict):
        return metadata

    if isinstance(metadata, str) and metadata.strip():

        try:
            parsed = json.loads(metadata)

            if isinstance(parsed, dict):
                return parsed

        except Exception:
            pass

    return {}


def _first_nonempty(*values):
    """
    Return first usable non-empty value.
    """

    for value in values:

        if value is None:
            continue

        value = str(value).strip()

        if (
            value
            and value.lower()
            not in {
                "nan",
                "none"
            }
        ):
            return value

    return ""


def _source_identity(
    row,
    metadata,
    source_type
):
    """
    Identify the original PQ or uploaded document.

    IMPORTANT:
    Chunk IDs must NOT be used as citation IDs.

    Every chunk belonging to the same document_id will
    therefore be grouped under one DOC citation.
    """

    if source_type == "Parliamentary Question":

        return _first_nonempty(
            row.get(
                "pq_id",
                ""
            ),
            metadata.get(
                "pq_id",
                ""
            ),
            row.get(
                "source_id",
                ""
            ),
            metadata.get(
                "source_id",
                ""
            ),
            metadata.get(
                "id",
                ""
            ),
            row.get(
                "title",
                ""
            ),
        )

    return _first_nonempty(
        row.get(
            "document_id",
            ""
        ),
        metadata.get(
            "document_id",
            ""
        ),
        row.get(
            "source_id",
            ""
        ),
        metadata.get(
            "source_id",
            ""
        ),
        row.get(
            "filename",
            ""
        ),
        metadata.get(
            "filename",
            ""
        ),
        row.get(
            "filepath",
            ""
        ),
        metadata.get(
            "filepath",
            ""
        ),
        row.get(
            "title",
            ""
        ),
    )


def _group_search_results(
    results,
    source_type,
    label_prefix,
    max_sources
):
    """
    Group semantic-search chunks by the original source.

    Example:

        Document A chunk 1
        Document A chunk 4
        Document A chunk 7

    becomes:

        [DOC1] Document A

    All three relevant passages are still supplied to the AI,
    but they share one citation.
    """

    if results is None or results.empty:
        return []

    grouped = {}
    group_order = []

    for row_number, (_, row) in enumerate(
        results.iterrows(),
        start=1
    ):

        metadata = _parse_source_metadata(
            row.get(
                "metadata",
                ""
            )
        )

        source_id = _source_identity(
            row,
            metadata,
            source_type
        )

        # ----------------------------------------------------
        # Emergency fallback
        #
        # Normally document_id / pq_id should always exist.
        # ----------------------------------------------------

        if not source_id:

            source_id = (
                f"__unknown_"
                f"{source_type}_"
                f"{row_number}"
            )

        # ----------------------------------------------------
        # Determine human-readable title
        # ----------------------------------------------------

        if source_type == "Parliamentary Question":

            title = _first_nonempty(
                row.get(
                    "title",
                    ""
                ),
                metadata.get(
                    "title",
                    ""
                ),
                metadata.get(
                    "question",
                    ""
                ),
                "Parliamentary Question"
            )

        else:

            title = _first_nonempty(
                row.get(
                    "title",
                    ""
                ),
                metadata.get(
                    "document_title",
                    ""
                ),
                metadata.get(
                    "title",
                    ""
                ),
                metadata.get(
                    "filename",
                    ""
                ),
                row.get(
                    "filename",
                    ""
                ),
                "NRF Document"
            )

        text = _first_nonempty(
            row.get(
                "text",
                ""
            )
        )

        try:
            similarity = float(
                row.get(
                    "similarity",
                    0
                ) or 0
            )

        except (
            TypeError,
            ValueError
        ):
            similarity = 0.0

        # ----------------------------------------------------
        # Create source group
        # ----------------------------------------------------

        if source_id not in grouped:

            grouped[source_id] = {

                "source_id":
                    source_id,

                "source_type":
                    source_type,

                "title":
                    title,

                "passages":
                    [],

                "similarity":
                    similarity,

                "metadata":
                    metadata,
            }

            group_order.append(
                source_id
            )

        group = grouped[
            source_id
        ]

        # Keep highest similarity for the source.
        group["similarity"] = max(
            group.get(
                "similarity",
                0
            ),
            similarity
        )

        # Preserve metadata.
        if (
            metadata
            and not group.get(
                "metadata"
            )
        ):
            group["metadata"] = metadata

        # Avoid duplicate passages.
        if (
            text
            and text
            not in group["passages"]
        ):
            group[
                "passages"
            ].append(text)

    # --------------------------------------------------------
    # Rank UNIQUE sources by their strongest matching chunk
    # --------------------------------------------------------

    ordered_groups = sorted(
        (
            grouped[source_id]
            for source_id
            in group_order
        ),
        key=lambda item:
            item.get(
                "similarity",
                0
            ),
        reverse=True
    )

    ordered_groups = ordered_groups[
        :max_sources
    ]

    sources = []

    # --------------------------------------------------------
    # Assign citation numbers AFTER grouping
    # --------------------------------------------------------

    for citation_number, group in enumerate(
        ordered_groups,
        start=1
    ):

        passages = group.pop(
            "passages",
            []
        )

        # All retrieved chunks remain available to the model.
        combined_text = "\n\n".join(
            (
                f"Relevant passage {i}:\n"
                f"{passage}"
            )
            for i, passage
            in enumerate(
                passages,
                start=1
            )
        )

        group["label"] = (
            f"{label_prefix}"
            f"{citation_number}"
        )

        group["text"] = (
            combined_text
        )

        group[
            "passage_count"
        ] = len(passages)

        sources.append(
            group
        )

    return sources


# ============================================================
# 🤖 CHATBOT SEARCH
# ============================================================

def retrieve_chatbot_sources(
    query,
    pq_top_k=5,
    document_top_k=5,
    min_similarity=0.20
):
    """
    Search Parliamentary Questions and NRF documents.

    Retrieval happens at CHUNK level.

    Citation numbering happens at SOURCE level.

    Therefore:

        Document A chunk 1
        Document A chunk 2
        Document A chunk 5

    all become:

        [DOC1]

    rather than:

        [DOC1]
        [DOC2]
        [DOC3]
    """

    # --------------------------------------------------------
    # Search MORE chunks than final unique source count.
    #
    # Example:
    # Top 5 chunks could all belong to one PDF.
    # Searching 20 chunks gives us a better chance of finding
    # several unique relevant documents.
    # --------------------------------------------------------

    pq_chunk_top_k = max(
        pq_top_k * 4,
        pq_top_k
    )

    document_chunk_top_k = max(
        document_top_k * 4,
        document_top_k
    )

    # --------------------------------------------------------
    # PQ retrieval
    # --------------------------------------------------------

    try:

        pq_results = (
            PQAI.search_parliamentary_questions(
                query=query,
                top_k=pq_chunk_top_k,
                min_similarity=min_similarity
            )
        )

    except Exception as exc:

        logger.exception(
            "PQ retrieval failed: %s",
            exc
        )

        pq_results = pd.DataFrame()

    # --------------------------------------------------------
    # Document retrieval
    # --------------------------------------------------------

    try:

        document_results = (
            PQAI.search_nrf_documents(
                query=query,
                top_k=document_chunk_top_k,
                min_similarity=min_similarity
            )
        )

    except Exception as exc:

        logger.exception(
            "Document retrieval failed: %s",
            exc
        )

        document_results = pd.DataFrame()

    # --------------------------------------------------------
    # GROUP PQ CHUNKS
    # --------------------------------------------------------

    pq_sources = _group_search_results(
        results=pq_results,
        source_type="Parliamentary Question",
        label_prefix="PQ",
        max_sources=pq_top_k
    )

    # --------------------------------------------------------
    # GROUP DOCUMENT CHUNKS
    # --------------------------------------------------------

    document_sources = _group_search_results(
        results=document_results,
        source_type="NRF Document",
        label_prefix="DOC",
        max_sources=document_top_k
    )

    return (
        pq_sources
        + document_sources
    )


# ============================================================
# 🤖 SOURCE CONTEXT BUILDER
# ============================================================

def build_chatbot_context(
    sources
):
    """
    Build source-labelled evidence for the AI.

    Each SOURCE has one citation label even if it contains
    several retrieved passages.
    """

    context_parts = []

    for source in sources:

        label = source.get(
            "label",
            ""
        )

        source_type = source.get(
            "source_type",
            ""
        )

        title = source.get(
            "title",
            ""
        )

        text = source.get(
            "text",
            ""
        )

        metadata = source.get(
            "metadata",
            {}
        )

        similarity = source.get(
            "similarity",
            0
        )

        passage_count = source.get(
            "passage_count",
            1
        )

        metadata = _parse_source_metadata(
            metadata
        )

        context_parts.append(
            f"""
SOURCE [{label}]
Source type: {source_type}
Title: {title}
Number of retrieved passages from this source: {passage_count}
Retrieval relevance: {similarity:.1%}
Metadata: {json.dumps(metadata, ensure_ascii=False)}

Relevant evidence from [{label}]:

{text}
""".strip()
        )

    return "\n\n---\n\n".join(
        context_parts
    )


# ============================================================
# 🤖 CHAT HISTORY BUILDER
# ============================================================

def build_recent_chat_history(
    messages,
    max_messages=8
):

    if not messages:
        return ""

    recent_messages = messages[
        -max_messages:
    ]

    parts = []

    for message in recent_messages:

        role = message.get(
            "role",
            ""
        )

        content = message.get(
            "content",
            ""
        )

        if role == "user":
            speaker = "User"

        elif role == "assistant":
            speaker = "NRF AI Assistant"

        else:
            continue

        parts.append(
            f"{speaker}:\n{content}"
        )

    return "\n\n".join(parts)


# ============================================================
# 🤖 GENERATE GROUNDED AI RESPONSE
# ============================================================

def generate_chatbot_answer(
    query,
    sources,
    chat_history
):

    client = get_openai_client()

    if client is None:
        raise RuntimeError(
            "OPENAI_API_KEY has not been configured."
        )

    if not sources:

        return (
            "I could not find sufficiently relevant information "
            "in the Parliamentary Questions or uploaded NRF "
            "documents to answer this question reliably."
        )

    context = build_chatbot_context(
        sources
    )

    recent_history = build_recent_chat_history(
        chat_history
    )

    valid_labels = [
        source.get(
            "label",
            ""
        )
        for source in sources
        if source.get(
            "label"
        )
    ]

    valid_labels_text = ", ".join(
        f"[{label}]"
        for label in valid_labels
    )

    instructions = """
You are the NRF Knowledge AI Assistant.

Your task is to answer questions using ONLY evidence supplied
from the NRF knowledge repository.

The repository contains:

1. Historical Parliamentary Questions and NRF answers.
2. NRF-related uploaded documents.

IMPORTANT SOURCE AND CITATION RULES:

- Ground factual claims in the supplied evidence.
- Do not invent NRF positions, policies, programmes, statistics,
  dates, commitments or explanations.
- If the evidence is insufficient, say that the repository does
  not contain enough information to answer confidently.

- Each [DOC#] represents ONE UNIQUE uploaded NRF document.
- Several relevant passages may appear underneath the same [DOC#].
- These passages are NOT separate documents.
- Never create a new citation number for an individual passage.

- Each [PQ#] represents ONE UNIQUE Parliamentary Question.
- Several passages belonging to the same PQ share that citation.

- Use ONLY the citation labels explicitly supplied in the
  retrieved evidence.
- Never invent a citation label.
- Never increment a citation number yourself.
- Never infer that passage 2 means DOC2.
- Cite the SOURCE label attached to that passage.

- Cite supporting evidence inline, for example:
  [PQ1]
  [DOC1]
  [PQ1][DOC2]

- Where useful, distinguish between evidence from Parliamentary
  Questions and evidence from uploaded NRF documents.

- Prefer synthesis over copying source text.
- Answer directly and clearly.
- Use concise headings or bullets only when they improve clarity.
- Do not mention semantic similarity scores.
"""

    prompt = f"""
CURRENT USER QUESTION

{query}


RECENT CONVERSATION

{recent_history or "No previous conversation."}


VALID CITATION LABELS

You may ONLY use these citation labels:

{valid_labels_text}


RETRIEVED NRF EVIDENCE

{context}


Answer the current user question using the retrieved NRF evidence.

Remember:

A citation identifies the UNIQUE SOURCE, not an individual
retrieved passage.

Do not create any citation label that is not in the VALID
CITATION LABELS list.
"""

    model = get_openai_model()

    response = client.responses.create(
        model=model,
        instructions=instructions,
        input=prompt
    )

    answer = getattr(
        response,
        "output_text",
        ""
    )

    if not answer:
        raise RuntimeError(
            "The AI model returned an empty response."
        )

    return answer.strip()



# ============================================================
# 🌐 SINGAPORE HANSARD — LIVE SEARCH / ANALYTICS
# ============================================================

HANSARD_BASE_URL = "https://sprs.parl.gov.sg"
HANSARD_SEARCH_URL = "https://sprs.parl.gov.sg/search/"
HANSARD_DISPLAY_ENDPOINT = "https://sprs.parl.gov.sg/search/getDisplayData"
HANSARD_TIMEOUT = 30
HANSARD_COLUMNS = [
    "date", "parliament", "title", "speaker", "content", "url", "source"
]


def _clean_hansard_text(value):
    if value is None:
        return ""
    text = BeautifulSoup(str(value), "html.parser").get_text(" ", strip=True)
    return re.sub(r"\s+", " ", text).strip()


def _first_hansard_value(record, *keys):
    if not isinstance(record, dict):
        return ""
    lowered = {str(k).lower(): v for k, v in record.items()}
    for key in keys:
        value = record.get(key)
        if value not in (None, "", [], {}):
            return value
        value = lowered.get(str(key).lower())
        if value not in (None, "", [], {}):
            return value
    return ""


def _extract_records_from_json(payload):
    candidates = []

    def walk(value):
        if isinstance(value, list):
            if value and all(isinstance(item, dict) for item in value):
                candidates.append(value)
            for item in value:
                walk(item)
        elif isinstance(value, dict):
            for child in value.values():
                walk(child)

    walk(payload)
    if not candidates:
        return []
    candidates.sort(key=len, reverse=True)
    return candidates[0]


def _normalise_hansard_record(record):
    title = _clean_hansard_text(_first_hansard_value(
        record, "title", "Title", "reportTitle", "report_title",
        "subject", "topic", "heading"
    ))
    content = _clean_hansard_text(_first_hansard_value(
        record, "content", "Content", "text", "Text", "snippet", "Snippet",
        "displayText", "description", "body"
    ))
    speaker = _clean_hansard_text(_first_hansard_value(
        record, "speaker", "Speaker", "member", "Member", "mp", "MP", "name"
    ))
    date_value = _first_hansard_value(
        record, "date", "Date", "sittingDate", "sitting_date",
        "reportDate", "report_date"
    )
    parliament = _clean_hansard_text(_first_hansard_value(
        record, "parliament", "Parliament", "parliamentNo",
        "parliamentNumber", "session"
    ))
    raw_url = str(_first_hansard_value(
        record, "url", "URL", "link", "Link", "href", "reportUrl", "reportURL"
    ) or "").strip()

    return {
        "date": _clean_hansard_text(date_value),
        "parliament": parliament,
        "title": title,
        "speaker": speaker,
        "content": content,
        "url": urljoin(HANSARD_BASE_URL, raw_url) if raw_url else "",
        "source": "Singapore Parliament Hansard",
    }


def _parse_hansard_html_results(html):
    soup = BeautifulSoup(html, "html.parser")
    records, seen = [], set()
    for link in soup.find_all("a", href=True):
        title = _clean_hansard_text(link.get_text(" ", strip=True))
        href = str(link.get("href", "")).strip()
        if not title:
            continue
        combined = f"{title.lower()} {href.lower()}"
        if not any(x in combined for x in ("report", "sitting", "sprs", "hansard")):
            continue
        url = urljoin(HANSARD_BASE_URL, href)
        identity = (url, title)
        if identity in seen:
            continue
        seen.add(identity)
        parent_text = link.parent.get_text(" ", strip=True) if link.parent else ""
        records.append({
            "date": "", "parliament": "", "title": title, "speaker": "",
            "content": parent_text, "url": url,
            "source": "Singapore Parliament Hansard",
        })
    return records


def search_hansard_live(
    query, start_date=None, end_date=None,
    search_mode="Any of the words", title_only=False, max_results=100
):
    query = str(query or "").strip()
    if not query:
        return pd.DataFrame(columns=HANSARD_COLUMNS), "Please enter a Hansard search query."

    option_map = {
        "All the words": "all", "Any of the words": "any",
        "Exact phrase": "exact", "Custom search": "custom",
    }
    option = option_map.get(search_mode, "any")
    start = start_date.strftime("%d/%m/%Y") if start_date else ""
    end = end_date.strftime("%d/%m/%Y") if end_date else ""

    payloads = [
        {
            "searchTerm": query, "keyword": query, "searchOption": option,
            "titleOnly": title_only, "startDate": start, "endDate": end,
            "page": 1, "pageSize": max_results,
        },
        {
            "keyword": query, "searchOption": option, "searchTitle": title_only,
            "fromDate": start, "toDate": end, "pageNo": 1,
            "pageSize": max_results,
        },
    ]
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120 Safari/537.36",
        "Accept": "application/json, text/plain, */*",
        "Referer": HANSARD_SEARCH_URL,
    }
    session = requests.Session()
    errors = []
    try:
        session.get(HANSARD_SEARCH_URL, headers=headers, timeout=HANSARD_TIMEOUT)
    except Exception as exc:
        logger.warning("Could not initialise Hansard session: %s", exc)

    for payload in payloads:
        for method in ("post", "get"):
            try:
                if method == "post":
                    response = session.post(
                        HANSARD_DISPLAY_ENDPOINT, json=payload,
                        headers=headers, timeout=HANSARD_TIMEOUT
                    )
                else:
                    response = session.get(
                        HANSARD_DISPLAY_ENDPOINT, params=payload,
                        headers=headers, timeout=HANSARD_TIMEOUT
                    )
                if not response.ok:
                    continue
                try:
                    raw = _extract_records_from_json(response.json())
                    records = [_normalise_hansard_record(r) for r in raw]
                except Exception:
                    records = _parse_hansard_html_results(response.text)

                records = [r for r in records if r.get("title") or r.get("content")]
                if records:
                    df = pd.DataFrame(records)
                    for col in HANSARD_COLUMNS:
                        if col not in df.columns:
                            df[col] = ""
                    df = (df[HANSARD_COLUMNS]
                          .drop_duplicates(subset=["date", "title", "url"])
                          .head(max_results).reset_index(drop=True))
                    return df, ""
            except Exception as exc:
                errors.append(f"{method.upper()}: {exc}")

    logger.warning("Hansard live search unsuccessful: %s", " | ".join(errors))
    return pd.DataFrame(columns=HANSARD_COLUMNS), (
        "The Singapore Parliament Hansard search endpoint did not return readable "
        "results. The Parliament website may have changed its internal search format."
    )


def prepare_hansard_dataframe(dataframe):
    if dataframe is None or dataframe.empty:
        return pd.DataFrame(columns=HANSARD_COLUMNS)
    df = dataframe.copy()
    for col in HANSARD_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df["date_parsed"] = pd.to_datetime(df["date"], errors="coerce", dayfirst=True)
    df["year"] = df["date_parsed"].dt.year
    df["month"] = df["date_parsed"].dt.to_period("M").astype(str)
    df["quarter"] = df["date_parsed"].dt.to_period("Q").astype(str)
    df["combined_text"] = (
        df["title"].fillna("").astype(str) + " " +
        df["content"].fillna("").astype(str)
    ).str.strip()
    return df


def hansard_yearly_counts(dataframe):
    df = prepare_hansard_dataframe(dataframe).dropna(subset=["year"])
    if df.empty:
        return pd.DataFrame(columns=["year", "records"])
    result = df.groupby("year").size().reset_index(name="records").sort_values("year")
    result["year"] = result["year"].astype(int)
    return result


def hansard_monthly_counts(dataframe):
    df = prepare_hansard_dataframe(dataframe).dropna(subset=["date_parsed"]).copy()
    if df.empty:
        return pd.DataFrame(columns=["month", "records"])
    df["month_period"] = df["date_parsed"].dt.to_period("M")
    result = (df.groupby("month_period").size().reset_index(name="records")
              .sort_values("month_period"))
    result["month"] = result["month_period"].dt.to_timestamp()
    return result[["month", "records"]]


def hansard_top_speakers(dataframe, top_n=15):
    if dataframe is None or dataframe.empty or "speaker" not in dataframe.columns:
        return pd.DataFrame(columns=["speaker", "records"])
    s = dataframe["speaker"].fillna("").astype(str).str.strip()
    s = s[s != ""]
    if s.empty:
        return pd.DataFrame(columns=["speaker", "records"])
    return s.value_counts().head(top_n).rename_axis("speaker").reset_index(name="records")


HANSARD_STOPWORDS = {
    "the", "and", "that", "this", "with", "from", "have", "has", "had", "for",
    "are", "was", "were", "will", "would", "could", "should", "there", "their",
    "they", "them", "into", "about", "which", "when", "where", "what", "who",
    "why", "how", "not", "but", "can", "our", "your", "you", "his", "her",
    "its", "also", "been", "being", "than", "then", "these", "those", "such",
    "more", "may", "parliament", "minister", "member", "members", "singapore",
}


def hansard_top_terms(dataframe, top_n=25):
    if dataframe is None or dataframe.empty:
        return pd.DataFrame(columns=["term", "count"])
    df = prepare_hansard_dataframe(dataframe)
    text = " ".join(df["combined_text"].fillna("").astype(str).tolist()).lower()
    words = re.findall(r"\b[a-z][a-z\-]{2,}\b", text)
    words = [w for w in words if w not in HANSARD_STOPWORDS]
    return pd.DataFrame(Counter(words).most_common(top_n), columns=["term", "count"])


def forecast_hansard_volume(dataframe, periods=6):
    monthly = hansard_monthly_counts(dataframe)
    if len(monthly) < 4:
        return pd.DataFrame(), "At least four months of dated Hansard results are required."
    if LinearRegression is None:
        return pd.DataFrame(), "scikit-learn is not installed."

    full_range = pd.date_range(monthly["month"].min(), monthly["month"].max(), freq="MS")
    monthly = (monthly.set_index("month").reindex(full_range, fill_value=0)
               .rename_axis("month").reset_index())
    monthly["t"] = np.arange(len(monthly))
    X = monthly[["t"]].values
    y = monthly["records"].astype(float).values
    model = LinearRegression().fit(X, y)
    fitted = model.predict(X)
    residuals = y - fitted
    residual_std = float(np.std(residuals, ddof=1)) if len(residuals) > 2 else 0.0

    future_t = np.arange(len(monthly), len(monthly) + periods)
    predictions = np.maximum(model.predict(future_t.reshape(-1, 1)), 0)
    future_months = pd.date_range(
        monthly["month"].max() + pd.offsets.MonthBegin(1), periods=periods, freq="MS"
    )
    margin = 1.96 * residual_std
    forecast = pd.DataFrame({
        "month": future_months,
        "forecast": predictions,
        "lower_95": np.maximum(predictions - margin, 0),
        "upper_95": predictions + margin,
        "type": "Forecast",
    })
    history = monthly[["month", "records"]].copy()
    history["type"] = "Historical"
    forecast["records"] = np.nan
    return pd.concat([history, forecast], ignore_index=True, sort=False), ""


def generate_hansard_ai_analysis(question, hansard_df):
    client = get_openai_client()
    if client is None:
        raise RuntimeError("OPENAI_API_KEY has not been configured.")
    if hansard_df is None or hansard_df.empty:
        return "No Hansard results are currently available for analysis."

    source_parts = []
    for i, (_, row) in enumerate(hansard_df.head(100).iterrows(), start=1):
        source_parts.append(
            f"""[H{i}]
Date: {row.get("date", "")}
Parliament: {row.get("parliament", "")}
Title: {row.get("title", "")}
Speaker: {row.get("speaker", "")}
Source URL: {row.get("url", "")}

Text:
{row.get("content", "")}"""
        )

    instructions = """
You are analysing Singapore Parliament Hansard search results.
Use ONLY the Hansard evidence supplied in the prompt.
Do not invent parliamentary statements, dates, speakers, statistics, policy positions
or trends. Distinguish observations in the retrieved dataset from conclusions about
Parliament as a whole. Cite evidence with the supplied [H#] labels only.
Do not predict political behaviour, election outcomes, government decisions or
individual MPs' future positions. Statistical forecasts concern record/query volume only.
"""
    response = client.responses.create(
        model=get_openai_model(),
        instructions=instructions,
        input=f"""USER ANALYSIS QUESTION

{question}

HANSARD SEARCH RESULTS

{chr(10).join(source_parts)}

Answer using only these Hansard results."""
    )
    answer = getattr(response, "output_text", "")
    if not answer:
        raise RuntimeError("The AI model returned an empty response.")
    return answer.strip()


def hansard_csv_bytes(dataframe):
    return dataframe.to_csv(index=False).encode("utf-8-sig")


def hansard_excel_bytes(dataframe):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        dataframe.to_excel(writer, index=False, sheet_name="Hansard Results")
        yearly = hansard_yearly_counts(dataframe)
        if not yearly.empty:
            yearly.to_excel(writer, index=False, sheet_name="Yearly Analytics")
        terms = hansard_top_terms(dataframe)
        if not terms.empty:
            terms.to_excel(writer, index=False, sheet_name="Top Terms")
    buffer.seek(0)
    return buffer.getvalue()



# ============================================================
# 📚 LOAD REPOSITORIES
# ============================================================

pq_df = load_csv(
    PQ_METADATA_FILE,
    PQ_COLUMNS
)

document_df = load_csv(
    DOCUMENT_METADATA_FILE,
    DOCUMENT_COLUMNS
)


# ============================================================
# 🧭 MAIN NAVIGATION
# ============================================================

tab_chat, tab_hansard, tab_pq, tab_documents, tab_admin = st.tabs(
    [
        "🤖 NRF AI Chatbot",
        "🌐 Live Hansard",
        "🏛️ Parliamentary Questions",
        "📚 NRF Documents",
        "⚙️ Repository Management",
    ]
)


# ============================================================
# 🤖 TAB 1 — NRF AI CHATBOT
# ============================================================

with tab_chat:

    st.header(
        "🤖 NRF Knowledge AI Chatbot"
    )

    st.markdown(
        """
        Ask questions about NRF using the knowledge stored in
        this repository.

        The chatbot searches both:

        - **Historical Parliamentary Questions and NRF answers**
        - **Uploaded NRF documents**

        before generating its response.
        """
    )

    api_key = get_openai_api_key()

    if not api_key:

        st.warning(
            """
            OpenAI API access has not been configured.

            Add `OPENAI_API_KEY` to Streamlit Secrets to enable
            AI-generated answers.
            """
        )

    # --------------------------------------------------------
    # Search-index health
    # --------------------------------------------------------

    try:

        search_index = (
            PQAI.load_search_index()
        )

    except Exception:

        search_index = pd.DataFrame()

    if search_index.empty:

        st.info(
            """
            The semantic search index is currently empty.

            Add Parliamentary Questions or NRF documents, then
            go to **Repository Management → Rebuild Semantic
            Search Index**.
            """
        )

    # --------------------------------------------------------
    # Chat state
    # --------------------------------------------------------

    if "nrf_chat_messages" not in st.session_state:

        st.session_state[
            "nrf_chat_messages"
        ] = []

    if "nrf_chat_sources" not in st.session_state:

        st.session_state[
            "nrf_chat_sources"
        ] = {}

    # --------------------------------------------------------
    # Clear chat
    # --------------------------------------------------------

    control_col1, control_col2 = st.columns(
        [5, 1]
    )

    with control_col2:

        if st.button(
            "🗑️ Clear chat",
            use_container_width=True
        ):

            st.session_state[
                "nrf_chat_messages"
            ] = []

            st.session_state[
                "nrf_chat_sources"
            ] = {}

            st.rerun()

    # --------------------------------------------------------
    # Welcome
    # --------------------------------------------------------

    if not st.session_state[
        "nrf_chat_messages"
    ]:

        with st.chat_message(
            "assistant"
        ):

            st.markdown(
                """
                Hello. I can answer questions using the NRF
                knowledge repository.

                For example:

                - What has NRF said about AI research?
                - What Parliamentary Questions have been asked
                  about R&D funding?
                - What do our uploaded documents say about
                  research manpower?
                - Compare Parliamentary responses with the
                  relevant NRF documents.
                """
            )

    # --------------------------------------------------------
    # Existing chat messages
    # --------------------------------------------------------

    for message_index, message in enumerate(
        st.session_state[
            "nrf_chat_messages"
        ]
    ):

        role = message.get(
            "role",
            "assistant"
        )

        content = message.get(
            "content",
            ""
        )

        with st.chat_message(role):

            st.markdown(content)

            if role == "assistant":

                sources = (
                    st.session_state[
                        "nrf_chat_sources"
                    ].get(
                        str(message_index),
                        []
                    )
                )

                if sources:

                    with st.expander(
                        f"📚 Sources used ({len(sources)})"
                    ):

                        for source in sources:

                            label = source.get(
                                "label",
                                ""
                            )

                            source_type = source.get(
                                "source_type",
                                ""
                            )

                            title = source.get(
                                "title",
                                ""
                            )

                            text = source.get(
                                "text",
                                ""
                            )

                            similarity = source.get(
                                "similarity",
                                0
                            )

                            passage_count = source.get(
                                "passage_count",
                                1
                            )

                            if (
                                source_type
                                == "Parliamentary Question"
                            ):
                                icon = "🏛️"

                            else:
                                icon = "📄"

                            st.markdown(
                                f"### {icon} [{label}] {title}"
                            )

                            st.caption(
                                f"{passage_count} relevant "
                                f"passage(s) retrieved • "
                                f"Best relevance: "
                                f"{similarity:.1%}"
                            )

                            st.write(text)

                            st.markdown("---")

    # --------------------------------------------------------
    # New question
    # --------------------------------------------------------

    user_question = st.chat_input(
        "Ask a question about NRF..."
    )

    if user_question:

        st.session_state[
            "nrf_chat_messages"
        ].append(
            {
                "role": "user",
                "content": user_question
            }
        )

        with st.chat_message(
            "user"
        ):

            st.markdown(
                user_question
            )

        with st.chat_message(
            "assistant"
        ):

            with st.spinner(
                "Searching Parliamentary Questions and "
                "NRF documents..."
            ):

                sources = (
                    retrieve_chatbot_sources(
                        query=user_question,
                        pq_top_k=5,
                        document_top_k=5,
                        min_similarity=0.20
                    )
                )

            if not sources:

                answer = (
                    "I could not find sufficiently relevant "
                    "information in the Parliamentary Questions "
                    "or uploaded NRF documents to answer this "
                    "question reliably."
                )

            elif not api_key:

                answer = (
                    "I found relevant repository sources, but "
                    "the OpenAI API key has not been configured, "
                    "so I cannot generate the AI response yet."
                )

            else:

                try:

                    with st.spinner(
                        "Preparing NRF-grounded answer..."
                    ):

                        answer = (
                            generate_chatbot_answer(
                                query=user_question,
                                sources=sources,
                                chat_history=(
                                    st.session_state[
                                        "nrf_chat_messages"
                                    ]
                                )
                            )
                        )

                except Exception as exc:

                    logger.exception(
                        "AI chatbot generation failed: %s",
                        exc
                    )

                    answer = (
                        "I found relevant NRF repository "
                        "information, but the AI response "
                        "could not be generated.\n\n"
                        f"Error: `{exc}`"
                    )

            st.markdown(answer)

            # ------------------------------------------------
            # Sources
            # ------------------------------------------------

            if sources:

                with st.expander(
                    f"📚 Sources used ({len(sources)})"
                ):

                    for source in sources:

                        label = source.get(
                            "label",
                            ""
                        )

                        source_type = source.get(
                            "source_type",
                            ""
                        )

                        title = source.get(
                            "title",
                            ""
                        )

                        text = source.get(
                            "text",
                            ""
                        )

                        similarity = source.get(
                            "similarity",
                            0
                        )

                        passage_count = source.get(
                            "passage_count",
                            1
                        )

                        if (
                            source_type
                            == "Parliamentary Question"
                        ):
                            icon = "🏛️"

                        else:
                            icon = "📄"

                        st.markdown(
                            f"### {icon} [{label}] {title}"
                        )

                        st.caption(
                            f"{passage_count} relevant "
                            f"passage(s) retrieved • "
                            f"Best relevance: "
                            f"{similarity:.1%}"
                        )

                        st.write(text)

                        st.markdown("---")

        assistant_index = len(
            st.session_state[
                "nrf_chat_messages"
            ]
        )

        st.session_state[
            "nrf_chat_messages"
        ].append(
            {
                "role": "assistant",
                "content": answer
            }
        )

        st.session_state[
            "nrf_chat_sources"
        ][
            str(assistant_index)
        ] = sources



# ============================================================
# 🌐 TAB 2 — LIVE SINGAPORE HANSARD
# ============================================================

with tab_hansard:
    st.header("🌐 Singapore Hansard Live Search & Analytics")
    st.markdown("""
    Search the **Singapore Parliament Official Reports (Hansard)** separately
    from the NRF knowledge repository. Live results are not automatically added
    to the NRF semantic search index.
    """)

    st.link_button("🏛️ Open Official Singapore Hansard", HANSARD_SEARCH_URL)

    if "hansard_results" not in st.session_state:
        st.session_state["hansard_results"] = pd.DataFrame(columns=HANSARD_COLUMNS)
    if "hansard_query" not in st.session_state:
        st.session_state["hansard_query"] = ""

    st.subheader("🔎 Live Hansard Search")

    with st.form("hansard_search_form"):
        c1, c2 = st.columns([3, 1])
        with c1:
            hansard_query = st.text_input(
                "Search Hansard",
                placeholder="e.g. artificial intelligence, research funding, quantum"
            )
        with c2:
            hansard_max_results = st.selectbox(
                "Maximum results", [25, 50, 100, 200], index=2
            )

        c1, c2 = st.columns(2)
        with c1:
            hansard_search_mode = st.selectbox(
                "Search mode",
                ["Any of the words", "All the words", "Exact phrase", "Custom search"]
            )
        with c2:
            hansard_title_only = st.checkbox("Search within titles only")

        use_date_filter = st.checkbox("Limit search by date")
        hansard_start_date = hansard_end_date = None
        if use_date_filter:
            c1, c2 = st.columns(2)
            with c1:
                hansard_start_date = st.date_input(
                    "From date", value=datetime(2015, 1, 1).date()
                )
            with c2:
                hansard_end_date = st.date_input(
                    "To date", value=datetime.now().date()
                )

        hansard_search_submit = st.form_submit_button(
            "🔎 Search Live Hansard", type="primary", use_container_width=True
        )

    if hansard_search_submit:
        if not hansard_query.strip():
            st.warning("Enter a keyword or phrase to search.")
        elif (use_date_filter and hansard_start_date and hansard_end_date
              and hansard_start_date > hansard_end_date):
            st.error("The start date cannot be after the end date.")
        else:
            with st.spinner("Searching the Singapore Parliament Hansard database..."):
                live_results, search_error = search_hansard_live(
                    query=hansard_query,
                    start_date=hansard_start_date if use_date_filter else None,
                    end_date=hansard_end_date if use_date_filter else None,
                    search_mode=hansard_search_mode,
                    title_only=hansard_title_only,
                    max_results=hansard_max_results,
                )
            st.session_state["hansard_results"] = live_results
            st.session_state["hansard_query"] = hansard_query
            if search_error:
                st.warning(search_error)
            if live_results.empty:
                st.info("No readable Hansard results were returned for this search.")
            else:
                st.success(f"{len(live_results)} Hansard record(s) retrieved.")

    hansard_results = st.session_state["hansard_results"]

    if hansard_results is not None and not hansard_results.empty:
        prepared_hansard = prepare_hansard_dataframe(hansard_results)
        st.markdown("---")
        st.subheader("📚 Current Hansard Dataset")
        if st.session_state.get("hansard_query"):
            st.caption(f"Current search: {st.session_state['hansard_query']}")

        m1, m2, m3, m4 = st.columns(4)
        with m1:
            st.metric("Records", len(prepared_hansard))
        with m2:
            st.metric("Dated records", int(prepared_hansard["date_parsed"].notna().sum()))
        with m3:
            speakers = prepared_hansard["speaker"].fillna("").astype(str).str.strip()
            st.metric("Speakers", speakers[speakers != ""].nunique())
        with m4:
            years = prepared_hansard["year"].dropna()
            if years.empty:
                year_range = "—"
            else:
                lo, hi = int(years.min()), int(years.max())
                year_range = str(lo) if lo == hi else f"{lo}–{hi}"
            st.metric("Period", year_range)

        c1, c2 = st.columns(2)
        with c1:
            st.download_button(
                "⬇️ Download CSV",
                data=hansard_csv_bytes(hansard_results),
                file_name="singapore_hansard_results.csv",
                mime="text/csv",
                use_container_width=True,
            )
        with c2:
            try:
                st.download_button(
                    "⬇️ Download Excel",
                    data=hansard_excel_bytes(hansard_results),
                    file_name="singapore_hansard_analysis.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
            except Exception as exc:
                st.caption(f"Excel export unavailable: {exc}")

        result_tab, ai_tab, analytics_tab, forecast_tab = st.tabs(
            ["📄 Results", "🤖 AI Analysis", "📊 Analytics", "📈 Forecast"]
        )

        with result_tab:
            st.subheader("Hansard Search Results")
            display_cols = [
                c for c in ["date", "parliament", "speaker", "title"]
                if c in prepared_hansard.columns
            ]
            st.dataframe(
                prepared_hansard[display_cols],
                use_container_width=True,
                hide_index=True,
            )
            st.markdown("---")
            for index, row in prepared_hansard.iterrows():
                title = row.get("title", "") or "Hansard Record"
                date = row.get("date", "")
                label = f"🏛️ {date} — {title}" if date else f"🏛️ {title}"
                with st.expander(label):
                    if row.get("speaker", ""):
                        st.markdown(f"**Speaker:** {row.get('speaker', '')}")
                    if row.get("parliament", ""):
                        st.caption(f"Parliament: {row.get('parliament', '')}")
                    if row.get("content", ""):
                        st.write(row.get("content", ""))
                    if row.get("url", ""):
                        st.link_button(
                            "🔗 Open Official Hansard Record",
                            row.get("url", ""),
                            key=f"hansard_link_{index}",
                        )

        with ai_tab:
            st.subheader("🤖 Analyse Retrieved Hansard Records")
            hansard_ai_question = st.text_area(
                "Analysis question",
                placeholder="e.g. What themes have emerged in discussion of AI?",
                height=120,
            )
            if st.button("🤖 Analyse Hansard", type="primary", key="analyse_hansard"):
                if not hansard_ai_question.strip():
                    st.warning("Enter an analysis question.")
                elif not get_openai_api_key():
                    st.error("OPENAI_API_KEY has not been configured.")
                else:
                    try:
                        with st.spinner("Analysing retrieved Hansard records..."):
                            answer = generate_hansard_ai_analysis(
                                hansard_ai_question, prepared_hansard
                            )
                        st.markdown(answer)
                    except Exception as exc:
                        logger.exception("Hansard AI analysis failed: %s", exc)
                        st.error(f"Hansard analysis failed: {exc}")

        with analytics_tab:
            st.subheader("📊 Hansard Data Analytics")
            c1, c2 = st.columns(2)
            with c1:
                st.markdown("### Records by Year")
                yearly = hansard_yearly_counts(prepared_hansard)
                if yearly.empty:
                    st.info("The retrieved records do not contain enough usable dates.")
                else:
                    st.bar_chart(yearly, x="year", y="records")
            with c2:
                st.markdown("### Most Frequent Speakers")
                speaker_counts = hansard_top_speakers(prepared_hansard)
                if speaker_counts.empty:
                    st.info("Speaker information is not available in these results.")
                else:
                    st.bar_chart(speaker_counts, x="speaker", y="records")

            st.markdown("### Parliamentary Attention Over Time")
            monthly = hansard_monthly_counts(prepared_hansard)
            if monthly.empty:
                st.info("Monthly trend analysis requires dated Hansard records.")
            else:
                st.line_chart(monthly, x="month", y="records")

            st.markdown("### Most Frequent Terms")
            terms = hansard_top_terms(prepared_hansard)
            if terms.empty:
                st.info("No terms could be extracted.")
            else:
                st.bar_chart(terms, x="term", y="count")
                st.dataframe(terms, use_container_width=True, hide_index=True)

        with forecast_tab:
            st.subheader("📈 Hansard Volume Forecast")
            st.info(
                "This forecasts the future volume of Hansard records matching the "
                "current query. It does not predict political outcomes, government "
                "decisions, or individual MPs' future behaviour."
            )
            periods = st.slider(
                "Forecast horizon (months)", min_value=3, max_value=24, value=6
            )
            forecast_data, forecast_error = forecast_hansard_volume(
                prepared_hansard, periods=periods
            )
            if forecast_error:
                st.warning(forecast_error)
            elif not forecast_data.empty:
                historical = forecast_data[forecast_data["type"] == "Historical"].copy()
                projected = forecast_data[forecast_data["type"] == "Forecast"].copy()
                st.markdown("### Historical Search Volume")
                st.line_chart(historical, x="month", y="records")
                st.markdown("### Forecast")
                st.line_chart(
                    projected.set_index("month")[["forecast", "lower_95", "upper_95"]]
                )
                display = projected[
                    ["month", "forecast", "lower_95", "upper_95"]
                ].copy()
                display["month"] = display["month"].dt.strftime("%Y-%m")
                for c in ["forecast", "lower_95", "upper_95"]:
                    display[c] = display[c].round(1)
                st.dataframe(display, use_container_width=True, hide_index=True)
                st.download_button(
                    "⬇️ Download Forecast CSV",
                    data=display.to_csv(index=False).encode("utf-8-sig"),
                    file_name="hansard_volume_forecast.csv",
                    mime="text/csv",
                )
    else:
        st.info(
            "Search the live Hansard database above to create a dataset for "
            "analysis and forecasting."
        )



# ============================================================
# 🏛️ TAB 3 — PARLIAMENTARY QUESTIONS
# ============================================================

with tab_pq:

    st.header(
        "🏛️ Parliamentary Questions & Answers"
    )

    st.markdown(
        "Add, browse and manage historical NRF "
        "Parliamentary Questions."
    )

    # ========================================================
    # ADD PQ
    # ========================================================

    with st.expander(
        "➕ Add Parliamentary Question",
        expanded=False
    ):

        with st.form(
            "add_pq_form"
        ):

            col1, col2 = st.columns(2)

            with col1:

                pq_date = st.date_input(
                    "Date"
                )

                mp_name = st.text_input(
                    "MP Name"
                )

                constituency = st.text_input(
                    "Constituency"
                )

            with col2:

                sitting = st.text_input(
                    "Parliamentary Sitting / Session"
                )

                topic = st.text_input(
                    "Topic",
                    placeholder=(
                        "e.g. AI, R&D funding, "
                        "research manpower"
                    )
                )

                keywords = st.text_input(
                    "Keywords",
                    placeholder=(
                        "AI, R&D, research funding"
                    )
                )

            question = st.text_area(
                "Parliamentary Question",
                height=180
            )

            answer = st.text_area(
                "NRF Answer",
                height=250
            )

            source = st.text_input(
                "Source / Reference"
            )

            submitted = st.form_submit_button(
                "💾 Save Parliamentary Question",
                type="primary"
            )

            if submitted:

                if (
                    not question.strip()
                    or not answer.strip()
                ):

                    st.error(
                        "Both the Parliamentary Question "
                        "and NRF Answer are required."
                    )

                else:

                    duplicate = pq_df[
                        (
                            pq_df["date"].astype(str)
                            == str(pq_date)
                        )
                        &
                        (
                            pq_df["mp_name"]
                            .astype(str)
                            .str.lower()
                            == mp_name.strip().lower()
                        )
                        &
                        (
                            pq_df["question"]
                            .astype(str)
                            .str.strip()
                            == question.strip()
                        )
                    ]

                    if not duplicate.empty:

                        st.warning(
                            "This Parliamentary Question "
                            "already appears to exist."
                        )

                    else:

                        new_pq = {
                            "pq_id":
                                generate_id("PQ"),

                            "date":
                                str(pq_date),

                            "sitting":
                                sitting.strip(),

                            "mp_name":
                                mp_name.strip(),

                            "constituency":
                                constituency.strip(),

                            "question":
                                question.strip(),

                            "answer":
                                answer.strip(),

                            "topic":
                                topic.strip(),

                            "keywords":
                                keywords.strip(),

                            "source":
                                source.strip(),

                            "uploaded_at":
                                datetime.now().isoformat()
                        }

                        pq_df = pd.concat(
                            [
                                pq_df,
                                pd.DataFrame(
                                    [new_pq]
                                )
                            ],
                            ignore_index=True
                        )

                        save_csv(
                            pq_df,
                            PQ_METADATA_FILE
                        )

                        with st.spinner(
                            "Updating semantic search index..."
                        ):

                            index_ok, _ = (
                                rebuild_search_index_safe(
                                    pq_df,
                                    document_df
                                )
                            )

                        if index_ok:

                            st.success(
                                "Parliamentary Question saved "
                                "and added to semantic search."
                            )

                        else:

                            st.warning(
                                "PQ saved, but the semantic "
                                "search index could not be rebuilt."
                            )

                        st.rerun()

    # ========================================================
    # PQ FILTER
    # ========================================================

    st.markdown("---")

    if pq_df.empty:

        st.info(
            "No Parliamentary Questions have been added yet."
        )

    else:

        col1, col2 = st.columns(2)

        with col1:

            pq_keyword_search = st.text_input(
                "Filter Parliamentary Questions",
                placeholder=(
                    "Search MP, topic, question or answer..."
                )
            )

        with col2:

            topic_options = [
                "All Topics"
            ] + sorted(
                [
                    x
                    for x in pq_df["topic"]
                    .dropna()
                    .astype(str)
                    .unique()
                    if x.strip()
                ]
            )

            selected_topic = st.selectbox(
                "Topic",
                topic_options
            )

        filtered_pq = pq_df.copy()

        if pq_keyword_search.strip():

            q = pq_keyword_search.lower()

            mask = filtered_pq.apply(
                lambda row:
                    q in " ".join(
                        [
                            str(row.get("mp_name", "")),
                            str(row.get("question", "")),
                            str(row.get("answer", "")),
                            str(row.get("topic", "")),
                            str(row.get("keywords", ""))
                        ]
                    ).lower(),
                axis=1
            )

            filtered_pq = filtered_pq[
                mask
            ]

        if selected_topic != "All Topics":

            filtered_pq = filtered_pq[
                filtered_pq["topic"]
                == selected_topic
            ]

        filtered_pq = filtered_pq.sort_values(
            "date",
            ascending=False
        )

        st.markdown(
            f"**{len(filtered_pq)} "
            f"Parliamentary Questions found**"
        )

        for _, row in filtered_pq.iterrows():

            title = (
                f"{row['date']} — "
                f"{row['mp_name'] or 'Unknown MP'}"
            )

            with st.expander(
                f"🏛️ {title}"
            ):

                st.markdown(
                    f"**Topic:** "
                    f"{row['topic'] or 'Not specified'}"
                )

                if row["sitting"]:

                    st.caption(
                        f"Sitting: {row['sitting']}"
                    )

                if row["constituency"]:

                    st.caption(
                        f"Constituency: "
                        f"{row['constituency']}"
                    )

                st.markdown(
                    "### Parliamentary Question"
                )

                st.write(
                    row["question"]
                )

                st.markdown(
                    "### NRF Answer"
                )

                st.write(
                    row["answer"]
                )

                if row["keywords"]:

                    st.caption(
                        f"Keywords: {row['keywords']}"
                    )

                if row["source"]:

                    st.caption(
                        f"Source: {row['source']}"
                    )

                if st.button(
                    "🗑️ Delete",
                    key=f"delete_pq_{row['pq_id']}"
                ):

                    pq_df = pq_df[
                        pq_df["pq_id"]
                        != row["pq_id"]
                    ]

                    save_csv(
                        pq_df,
                        PQ_METADATA_FILE
                    )

                    rebuild_search_index_safe(
                        pq_df,
                        document_df
                    )

                    st.rerun()


# ============================================================
# 📚 TAB 4 — NRF DOCUMENTS
# ============================================================

with tab_documents:

    st.header(
        "📚 NRF Document Repository"
    )

    st.write(
        """
        Upload NRF-related PDF, Word, Excel and text documents.

        Each uploaded file is treated as ONE unique document for
        chatbot citations, even though it may be divided into
        several semantic-search chunks internally.
        """
    )

    uploaded_files = st.file_uploader(
        "Upload NRF Documents",
        type=[
            "pdf",
            "docx",
            "xlsx",
            "xls",
            "txt"
        ],
        accept_multiple_files=True,
        key="nrf_documents"
    )

    if uploaded_files:

        for uploaded_file in uploaded_files:

            st.markdown(
                f"### `{uploaded_file.name}`"
            )

            col1, col2 = st.columns(2)

            with col1:

                document_title = st.text_input(
                    "Document title",
                    value=os.path.splitext(
                        uploaded_file.name
                    )[0],
                    key=f"title_{uploaded_file.name}"
                )

                document_type = st.selectbox(
                    "Document type",
                    [
                        "Annual Report",
                        "Strategy / Policy",
                        "Research / R&D",
                        "Programme / Funding",
                        "Statistics / Dataset",
                        "Speech",
                        "Press Release",
                        "Parliamentary Material",
                        "Other"
                    ],
                    key=f"type_{uploaded_file.name}"
                )

            with col2:

                document_keywords = st.text_input(
                    "Keywords",
                    key=f"keywords_{uploaded_file.name}"
                )

                document_source = st.text_input(
                    "Source / Reference",
                    key=f"source_{uploaded_file.name}"
                )

            description = st.text_area(
                "Description",
                key=f"description_{uploaded_file.name}"
            )

            if st.button(
                f"📥 Add `{uploaded_file.name}` to Repository",
                key=f"add_doc_{uploaded_file.name}",
                type="primary"
            ):

                existing = document_df[
                    document_df["filename"]
                    .astype(str)
                    .str.lower()
                    == uploaded_file.name.lower()
                ]

                if not existing.empty:

                    st.warning(
                        f"`{uploaded_file.name}` is already "
                        "in the repository."
                    )

                    continue

                save_path = ensure_unique_path(
                    DOCUMENT_FOLDER,
                    uploaded_file.name
                )

                try:

                    with open(
                        save_path,
                        "wb"
                    ) as f:

                        f.write(
                            uploaded_file.getbuffer()
                        )

                except Exception as exc:

                    st.error(
                        f"Could not save "
                        f"`{uploaded_file.name}`: {exc}"
                    )

                    continue

                with st.spinner(
                    "Extracting document contents..."
                ):

                    extracted_text = extract_raw_text(
                        save_path
                    )

                if not extracted_text.strip():

                    st.error(
                        "No text could be extracted from "
                        "this document."
                    )

                    try:
                        os.remove(save_path)
                    except OSError:
                        pass

                    continue

                # IMPORTANT:
                # One permanent ID is created per uploaded file.
                document_id = generate_id(
                    "DOC"
                )

                new_document = {

                    "document_id":
                        document_id,

                    "filename":
                        uploaded_file.name,

                    "filepath":
                        save_path,

                    "document_title":
                        document_title.strip(),

                    "document_type":
                        document_type,

                    "description":
                        description.strip(),

                    "keywords":
                        document_keywords.strip(),

                    "source":
                        document_source.strip(),

                    "uploaded_at":
                        datetime.now().isoformat()
                }

                document_df = pd.concat(
                    [
                        document_df,
                        pd.DataFrame(
                            [new_document]
                        )
                    ],
                    ignore_index=True
                )

                save_csv(
                    document_df,
                    DOCUMENT_METADATA_FILE
                )

                with st.spinner(
                    "Adding document to semantic search index..."
                ):

                    index_ok, _ = (
                        rebuild_search_index_safe(
                            pq_df,
                            document_df
                        )
                    )

                if index_ok:

                    st.success(
                        f"`{uploaded_file.name}` successfully "
                        "added to the NRF repository."
                    )

                else:

                    st.warning(
                        f"`{uploaded_file.name}` was added, "
                        "but the search index rebuild failed."
                    )

    st.markdown("---")

    st.subheader(
        "📂 Uploaded NRF Documents"
    )

    document_df = load_csv(
        DOCUMENT_METADATA_FILE,
        DOCUMENT_COLUMNS
    )

    if document_df.empty:

        st.info(
            "No NRF documents have been uploaded."
        )

    else:

        document_df = document_df.sort_values(
            "uploaded_at",
            ascending=False
        )

        st.markdown(
            f"**{len(document_df)} "
            f"documents in repository**"
        )

        for _, row in document_df.iterrows():

            document_title_display = (
                row["document_title"]
                or row["filename"]
                or "Untitled Document"
            )

            with st.expander(
                f"📄 {document_title_display}"
            ):

                st.caption(
                    f"Document ID: "
                    f"{row['document_id']}"
                )

                st.caption(
                    f"Filename: {row['filename']}"
                )

                st.caption(
                    f"Type: {row['document_type']}"
                )

                if row["description"]:

                    st.write(
                        row["description"]
                    )

                if row["keywords"]:

                    st.caption(
                        f"Keywords: {row['keywords']}"
                    )

                if row["source"]:

                    st.caption(
                        f"Source: {row['source']}"
                    )

                col1, col2 = st.columns(2)

                with col1:

                    filepath = row[
                        "filepath"
                    ]

                    if (
                        isinstance(filepath, str)
                        and os.path.exists(filepath)
                    ):

                        try:

                            with open(
                                filepath,
                                "rb"
                            ) as f:

                                file_bytes = f.read()

                            st.download_button(
                                "⬇️ Download",
                                data=file_bytes,
                                file_name=row[
                                    "filename"
                                ],
                                key=(
                                    f"download_"
                                    f"{row['document_id']}"
                                )
                            )

                        except Exception:

                            st.caption(
                                "⚠️ Stored file could not "
                                "be opened."
                            )

                with col2:

                    if st.button(
                        "🗑️ Delete",
                        key=(
                            f"delete_doc_"
                            f"{row['document_id']}"
                        )
                    ):

                        filepath = row[
                            "filepath"
                        ]

                        if (
                            isinstance(filepath, str)
                            and os.path.exists(filepath)
                        ):

                            try:
                                os.remove(filepath)
                            except Exception:
                                pass

                        document_df = document_df[
                            document_df[
                                "document_id"
                            ]
                            != row["document_id"]
                        ]

                        save_csv(
                            document_df,
                            DOCUMENT_METADATA_FILE
                        )

                        rebuild_search_index_safe(
                            pq_df,
                            document_df
                        )

                        st.rerun()


# ============================================================
# ⚙️ TAB 5 — REPOSITORY MANAGEMENT
# ============================================================

with tab_admin:

    st.header(
        "⚙️ Repository Management"
    )

    pq_df = load_csv(
        PQ_METADATA_FILE,
        PQ_COLUMNS
    )

    document_df = load_csv(
        DOCUMENT_METADATA_FILE,
        DOCUMENT_COLUMNS
    )

    try:

        search_index = (
            PQAI.load_search_index()
        )

    except Exception as exc:

        logger.warning(
            "Could not load search index: %s",
            exc
        )

        search_index = pd.DataFrame()

    col1, col2, col3 = st.columns(3)

    with col1:

        st.metric(
            "Parliamentary Questions",
            len(pq_df)
        )

    with col2:

        st.metric(
            "NRF Documents",
            len(document_df)
        )

    with col3:

        st.metric(
            "Searchable Chunks",
            len(search_index)
        )

    st.markdown("---")

    st.subheader(
        "🔄 Search Index"
    )

    if st.button(
        "🔄 Rebuild Semantic Search Index",
        type="primary"
    ):

        with st.spinner(
            "Rebuilding NRF semantic search index..."
        ):

            index_ok, rebuilt = (
                rebuild_search_index_safe(
                    pq_df,
                    document_df
                )
            )

        if index_ok:

            st.success(
                f"Search index rebuilt successfully. "
                f"{len(rebuilt)} searchable chunks created."
            )

        else:

            st.error(
                "Search index rebuild failed."
            )

    st.markdown("---")

    st.subheader(
        "🩺 Search Index Health"
    )

    embeddings = None

    try:

        embeddings = (
            PQAI.load_embeddings()
        )

    except Exception as exc:

        logger.warning(
            "Could not load embeddings: %s",
            exc
        )

    index_count = len(
        search_index
    )

    if embeddings is None:

        embedding_count = 0

    else:

        try:
            embedding_count = len(
                embeddings
            )

        except TypeError:
            embedding_count = 0

    health_col1, health_col2 = st.columns(
        2
    )

    with health_col1:

        st.metric(
            "Indexed passages",
            index_count
        )

    with health_col2:

        st.metric(
            "Embedding vectors",
            embedding_count
        )

    if (
        index_count
        == embedding_count
        and index_count > 0
    ):

        st.success(
            "Search index and embeddings are synchronized."
        )

    elif (
        index_count == 0
        and embedding_count == 0
    ):

        st.info(
            "The search index is currently empty."
        )

    else:

        st.warning(
            "The search index and embeddings are not "
            "synchronized. Rebuild the semantic search index."
        )

    st.markdown("---")

    st.subheader(
        "📊 Repository Data"
    )

    st.write(
        "Parliamentary Questions"
    )

    if not pq_df.empty:

        st.dataframe(
            pq_df,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.caption(
            "No Parliamentary Questions in repository."
        )

    st.write(
        "NRF Documents"
    )

    if not document_df.empty:

        st.dataframe(
            document_df,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.caption(
            "No NRF documents in repository."
        )